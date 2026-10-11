import os
import secrets
import shutil
import json
import re
import math
import uuid
from datetime import datetime, date
from typing import List, Optional
from contextlib import asynccontextmanager
from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Header, Depends
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from backend.db import get_db_connection, init_db, get_db_path, get_db_dir
import backend.calculations as calc
from backend.formatting import inr
from backend.risk import EQUITY_ORIENTED_FUND_TYPES, effective_risk, fund_type, metadata as risk_metadata
import backend.backups as backups
import backend.tax.engine as tax_engine
import backend.tax.rules_loader as tax_rules
import backend.tax.rule_store as tax_store
import backend.tax.scenario_store as tax_scenarios
import backend.tax.nps as tax_nps
import backend.tax.payroll as tax_payroll
import backend.tax.waterfall as tax_waterfall
import backend.tax.sensitivity as tax_sensitivity_mod
import backend.tax.epf as tax_epf
import backend.tax.prefill as tax_prefill_mod
import backend.tax.payslip_pdf as tax_payslip_pdf_mod
import backend.tax.advance_tax as tax_advance_mod
import backend.tax.harvesting as tax_harvesting_mod
from backend.decision_engine import generate_recommendations, generate_credit_card_annual_reviews

# Shared-secret auth for the Docker/web deployment path only. Empty/unset (the
# default, and the only mode desktop/Electron ever uses) disables this check
# entirely — zero-config desktop startup is unaffected. This is deliberately
# NOT a multi-user login system: one operator-set secret, checked on every
# request, no sessions/passwords/accounts. See
# Architecture/11-Security-Architecture.md for the full
# threat model and what this does/doesn't protect against (no rate limiting,
# no TLS — a reverse proxy should terminate TLS if this ever crosses a network).
API_SECRET = os.environ.get("ARTHA_API_SECRET", "").strip()

def require_api_secret(x_artha_api_secret: str = Header(default="")):
    if not API_SECRET:
        return
    if not secrets.compare_digest(x_artha_api_secret, API_SECRET):
        raise HTTPException(status_code=401, detail="Invalid or missing API secret")

@asynccontextmanager
async def lifespan(_app):
    # Back up first (every platform: desktop, web, Docker), so a migration can never be the
    # only copy of the data; skipped when nothing changed since the last backup.
    try:
        backups.auto_backup("open")
    except Exception as e:  # a failed backup must not stop the app from starting
        print("[Backup] Startup backup failed:", e)
    # Schema/migrations run once at startup, before the first request.
    init_db()
    yield

app = FastAPI(title="Artha API", dependencies=[Depends(require_api_secret)], lifespan=lifespan)

def _parse_allowed_origins() -> list:
    raw = os.environ.get("ARTHA_ALLOWED_ORIGINS", "").strip()
    if raw:
        return [o.strip() for o in raw.split(",") if o.strip()]
    # Safe default matching today's real desktop/dev usage (Vite dev server
    # only) — never silently falls back to "*". Derived from ARTHA_FRONTEND_PORT
    # (same env var electron/main.js and vite.config.mts already read, default
    # 5173) so this stays correct for any dev setup that overrides the frontend
    # port — e.g. start.sh's 5183 default — without also having to set
    # ARTHA_ALLOWED_ORIGINS by hand. Docker/web deployments that need a
    # different browser-facing origin (a real hostname, HTTPS, etc.) must still
    # set ARTHA_ALLOWED_ORIGINS explicitly. Packaged Electron loads the
    # frontend via file:// and either sends no Origin header or "null", neither
    # of which needs (or should be) listed here — CORS is a browser-enforced,
    # not server-enforced, concept and doesn't apply to that request shape.
    frontend_port = os.environ.get("ARTHA_FRONTEND_PORT", "5173").strip() or "5173"
    return [
        f"http://localhost:{frontend_port}",
        f"http://127.0.0.1:{frontend_port}",
    ]

# Enable CORS for frontend development
app.add_middleware(
    CORSMiddleware,
    allow_origins=_parse_allowed_origins(),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def _with_metadata(row: dict) -> dict:
    """Custom fields (category fields_definition) flattened onto the row, as for assets and
    liabilities; a custom key never overrides a real column."""
    try:
        meta = json.loads(row.get("metadata_json") or "{}")
    except ValueError:
        meta = {}
    if isinstance(meta, dict):
        for k, v in meta.items():
            row.setdefault(k, v)
    return row

def default_owner_id(cursor) -> int:
    """Owner for a record created without one (e.g. from the Credit Cards screen): the
    Primary Owner, else the first person, else 1 for an empty people table."""
    row = cursor.execute("""SELECT id FROM people ORDER BY CASE role WHEN 'Primary Owner' THEN 0 ELSE 1 END, id LIMIT 1""").fetchone()
    return row["id"] if row else 1

def get_table_columns(cursor, table_name: str) -> list:
    cursor.execute(f"PRAGMA table_info({table_name})")
    return [row["name"] for row in cursor.fetchall()]

def serialize_entity_payload(cursor, table_name: str, payload: dict) -> dict:
    cols = get_table_columns(cursor, table_name)
    standard_params = {}
    extra_params = {}
    
    for k, v in payload.items():
        if k in cols:
            standard_params[k] = v
        else:
            extra_params[k] = v
            
    if "metadata_json" in cols:
        standard_params["metadata_json"] = json.dumps(extra_params)
    elif "funding_sources_json" in cols:
        standard_params["funding_sources_json"] = json.dumps(extra_params)
    elif "snapshot_data" in cols:
        standard_params["snapshot_data"] = json.dumps(extra_params)
        
    return standard_params

# --- Pydantic Models for Requests ---
class AccountSchema(BaseModel):
    name: str
    institution: str
    current_balance: Optional[float] = 0.0
    remarks: Optional[str] = ""

class LoanSchema(BaseModel):
    name: str
    lender: str
    outstanding: float
    roi: float
    emi: float
    tenure: int # original tenure in months
    start_date: str # YYYY-MM-DD
    prepayment_allowed: bool
    processing_fee: float
    loan_type: str

class InvestmentSchema(BaseModel):
    name: str
    type: str # Mutual Fund, Stock, PPF, EPF, Gold, NPS, etc.
    current_value: float
    invested_amount: float
    goal_mapping: Optional[str] = ""

class PropertySchema(BaseModel):
    name: str
    purchase_price: float
    current_value: float
    associated_loan_id: Optional[int] = None
    rental_income: float = 0.0
    expenses: float = 0.0
    remarks: Optional[str] = ""

class InsuranceSchema(BaseModel):
    company: str
    policy_number: str
    premium: float
    frequency: str # Monthly, Quarterly, Half-Yearly, Yearly
    coverage: float
    nominee: str
    maturity_date: Optional[str] = ""
    last_premium_date: Optional[str] = ""
    renewal_date: Optional[str] = ""
    waiting_period: Optional[str] = ""
    claim_history: Optional[str] = ""
    insurance_type: Optional[str] = "Traditional Life"
    surrender_value: Optional[float] = 0.0

class GoalSchema(BaseModel):
    name: str
    target_amount: float
    current_amount: float
    monthly_contribution: float
    goal_type: str

class RuleSchema(BaseModel):
    rule_key: str
    value: float
    is_enabled: bool

class CashFlowSchema(BaseModel):
    id: Optional[int] = None
    month: str
    net_salary: float
    other_income: float
    household_expenses: float
    emi: float
    insurance_premium: float
    investments: float
    one_time_expenses: float
    notes: Optional[str] = ""

class CashFlowSimulationRequest(BaseModel):
    net_salary: float
    other_income: float
    household_expenses: float
    emi: float
    insurance_premium: float
    investments: float
    one_time_expenses: float
    
    # Adjustments
    salary_change_pct: Optional[float] = 0.0
    salary_change_abs: Optional[float] = 0.0
    one_time_bonus: Optional[float] = 0.0
    new_emi: Optional[float] = 0.0
    emi_reduction: Optional[float] = 0.0
    additional_investment: Optional[float] = 0.0
    expense_reduction_pct: Optional[float] = 0.0
    expense_reduction_abs: Optional[float] = 0.0

class SimulationRequestSchema(BaseModel):
    scenario_type: str # close_loan, delay_goal, increase_sip, salary_change, receive_bonus, roi_reduction, brother_property
    params: dict # parameters specific to the scenario

class DebtPayoffRequest(BaseModel):
    extra_monthly: float = 0.0
    strategy: str = "avalanche"
    prepayment_frequency: str = "monthly"
    prepayment_amount: float = 0.0
    prepayment_month: int = 3

class NetWorthForecastRequest(BaseModel):
    horizon_years: int = 10
    expected_return: float = 12.0
    monthly_savings: float = 20000.0
    inflation_rate: float = 6.0

# --- Core Dynamic Aggregate / Dashboard Data ---
@app.get("/api/init")
def get_init_data():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        # Accounts (derived balances computed on the fly)
        cursor.execute("""
            SELECT
                a.*,
                i.name as institution_name,
                p.name as owner_name,
                c.name as category_name,
                (
                    SELECT COALESCE(SUM(t.amount), 0.0)
                    FROM transactions t
                    WHERE t.dest_ref_type = 'Account' AND t.dest_ref_id = a.id
                ) - (
                    SELECT COALESCE(SUM(t.amount), 0.0)
                    FROM transactions t
                    WHERE t.source_ref_type = 'Account' AND t.source_ref_id = a.id
                ) AS current_balance,
                (
                    SELECT COALESCE(SUM(t.amount), 0.0)
                    FROM transactions t
                    WHERE t.dest_ref_type = 'Account' AND t.dest_ref_id = a.id
                      AND t.notes IN ('Initial balance registration', 'CSV Import Initial Balance')
                ) AS opening_balance
            FROM accounts a
            LEFT JOIN institutions i ON a.institution_id = i.id
            LEFT JOIN people p ON a.owner_id = p.id
            LEFT JOIN categories c ON a.category_id = c.id
        """)
        accounts = []
        for row in cursor.fetchall():
            d = dict(row)
            if d.get("metadata_json"):
                try:
                    meta = json.loads(d["metadata_json"])
                    d.update(meta)
                except:
                    pass
            accounts.append(d)
        
        # Assets (Mutual funds, stocks, gold, EPFO, property, insurance etc.)
        cursor.execute("""
            SELECT a.*, c.name as category_name, i.name as institution_name, p.name as owner_name
            FROM assets a
            LEFT JOIN categories c ON a.category_id = c.id
            LEFT JOIN institutions i ON a.institution_id = i.id
            LEFT JOIN people p ON a.owner_id = p.id
        """)
        assets = []
        for r in cursor.fetchall():
            d = dict(r)
            if d.get("metadata_json"):
                try:
                    d.update(json.loads(d["metadata_json"]))
                except:
                    pass
            d["effective_risk"] = effective_risk(d.get("category_name"), d.get("risk_level"), d.get("metadata_json"))
            if d.get("category_name") == "Insurance":
                cursor.execute("SELECT * FROM insurance_details WHERE asset_id = ?", (d["id"],))
                details_row = cursor.fetchone()
                if details_row:
                    d["insurance_details"] = dict(details_row)
                    d.update({k: v for k, v in dict(details_row).items() if k != "asset_id"})
                cursor.execute("SELECT * FROM insurance_claims WHERE asset_id = ? ORDER BY claim_date", (d["id"],))
                d["insurance_claims"] = [dict(row) for row in cursor.fetchall()]
                cursor.execute("SELECT * FROM insurance_premium_payments WHERE asset_id = ? ORDER BY payment_date", (d["id"],))
                d["premium_payments"] = [dict(row) for row in cursor.fetchall()]
                # Cost basis for the Assets grid: sum of logged premium payments, falling
                # back to invested_amount (where legacy/manual entries recorded premiums
                # paid so far) when no payments have been logged yet.
                paid = sum(p["amount"] or 0.0 for p in d["premium_payments"])
                d["premiums_paid"] = paid if paid > 0 else (d.get("invested_amount") or 0.0)
                # Unified coverage figure across policy types: Life/Term use sum_assured,
                # Health uses sum_insured, Vehicle uses idv (matches the Dashboard's
                # Insurance Coverage panel logic).
                d["cover"] = max(
                    d.get("sum_assured") or 0.0,
                    d.get("sum_insured") or 0.0,
                    d.get("idv") or 0.0,
                )
            if d.get("category_name") == "Property":
                cursor.execute("SELECT * FROM property_rental_ledger WHERE asset_id = ? ORDER BY month", (d["id"],))
                ledger = [dict(row) for row in cursor.fetchall()]
                d["rental_ledger"] = ledger
                if ledger:
                    latest = ledger[-1]
                    d["rental_income"] = latest["rent_amount"]
                    d["maintenance_expenses"] = latest["expenses_amount"]
            if d.get("category_name") in ("Stock", "Mutual Fund"):
                cursor.execute("SELECT * FROM holding_lots WHERE asset_id = ? ORDER BY purchase_date", (d["id"],))
                lots = [dict(lot) for lot in cursor.fetchall()]
                d["holding_lots"] = lots
                # No recorded cost (e.g. the auto-migrated 1-unit, ₹0 opening lot): a return
                # on nothing paid is meaningless, so leave xirr unset.
                if lots and sum(lot["quantity"] * lot["price_per_unit"] for lot in lots) > 0:
                    cash_flows = [
                        {"date": lot["purchase_date"], "amount": -(lot["quantity"] * lot["price_per_unit"])}
                        for lot in lots
                    ]
                    cash_flows.append({"date": date.today().strftime("%Y-%m-%d"), "amount": d.get("current_value", 0.0)})
                    d["xirr"] = round(calc.calculate_xirr(cash_flows), 2)
            assets.append(d)

        # Liabilities (Home loans, car loans, credit cards)
        cursor.execute("""
            SELECT l.*, c.name as category_name, i.name as lender_name, p.name as owner_name
            FROM liabilities l
            LEFT JOIN categories c ON l.category_id = c.id
            LEFT JOIN institutions i ON l.lender_institution_id = i.id
            LEFT JOIN people p ON l.owner_id = p.id
        """)
        liabilities_raw = cursor.fetchall()
        liabilities = []
        for r in liabilities_raw:
            d = dict(r)
            if d.get("metadata_json"):
                try:
                    d.update(json.loads(d["metadata_json"]))
                except:
                    pass
            if d.get("category_name") == "Credit Card":
                cursor.execute("SELECT * FROM credit_card_details WHERE liability_id = ?", (d["id"],))
                details_row = cursor.fetchone()
                if details_row:
                    d["credit_card_details"] = dict(details_row)
                    d.update({k: v for k, v in dict(details_row).items() if k != "liability_id"})
                real_spend = compute_real_cycle_spend_total(cursor, d["id"], d.get("statement_date"))
                d["real_cycle_spend_total"] = real_spend["total"]
                d["has_spend_transactions"] = real_spend["has_any_ever"]
                d["cycle_start"] = real_spend["cycle_start"]
                d["cycle_end"] = real_spend["cycle_end"]
                d["statement_spend_adjustment"] = real_spend["adjustment_amount"]
            # Add dynamic tenure and progress math if it is a reducing balance loan
            if d.get("outstanding", 0.0) > 0 and d.get("interest_rate", 0.0) > 0 and d.get("emi", 0.0) > 0:
                rem_n = calc.calculate_remaining_tenure(d["outstanding"], d["interest_rate"], d["emi"])
                d["remaining_tenure"] = round(rem_n, 1) if rem_n != float('inf') else -1
                d["remaining_interest"] = round(calc.calculate_remaining_interest(d["outstanding"], d["interest_rate"], d["emi"]), 2)
                d["progress"] = round(calc.calculate_loan_progress(d.get("tenure_months", 120), rem_n), 1)
            else:
                d["remaining_tenure"] = -1
                d["remaining_interest"] = 0.0
                d["progress"] = 100.0
            liabilities.append(d)

        # Income
        cursor.execute("""
            SELECT inc.*, c.name as category_name, p.name as owner_name
            FROM income inc
            LEFT JOIN categories c ON inc.category_id = c.id
            LEFT JOIN people p ON inc.owner_id = p.id
        """)
        income = [_with_metadata(dict(row)) for row in cursor.fetchall()]

        # Expenses
        cursor.execute("""
            SELECT e.*, c.name as category_name, p.name as owner_name
            FROM expenses e
            LEFT JOIN categories c ON e.category_id = c.id
            LEFT JOIN people p ON e.owner_id = p.id
        """)
        expenses = [_with_metadata(dict(row)) for row in cursor.fetchall()]

        # Goals: current_amount is derived from linked holdings when any exist, else the
        # manually-entered stored value (goals funded outside tracked assets stay manual).
        # The alias must NOT be named current_amount: g.* already carries that column, and
        # sqlite3.Row name lookup resolves duplicate names to the LEFTMOST column — a
        # same-named alias is silently shadowed by the stored value (the original bug).
        cursor.execute("""
            SELECT
                g.*,
                CASE WHEN EXISTS (SELECT 1 FROM goal_holdings gh WHERE gh.goal_id = g.id)
                    THEN COALESCE(
                        (
                            SELECT SUM(a.current_value * gh.allocation_weight)
                            FROM assets a
                            JOIN goal_holdings gh ON a.id = gh.asset_id
                            WHERE gh.goal_id = g.id
                        ), 0.0
                    )
                    ELSE g.current_amount
                END AS derived_current_amount
            FROM goals g
        """)
        goals_raw = cursor.fetchall()
        goals = []
        for r in goals_raw:
            d = dict(r)
            d["current_amount"] = d.pop("derived_current_amount")
            if d.get("funding_sources_json"):
                try:
                    d["funding_sources"] = json.loads(d["funding_sources_json"])
                except:
                    d["funding_sources"] = []
            # The goal form's "Only low-risk holdings" lives in that blob.
            d["low_risk_only"] = bool(isinstance(d.get("funding_sources"), dict) and d["funding_sources"].get("low_risk_only"))
            comp_details = calc.calculate_goal_expected_completion(d["target_amount"], d["current_amount"], d.get("monthly_contribution", 0.0))
            d.update(comp_details)
            goals.append(d)

        # Rules
        cursor.execute("SELECT * FROM financial_rules")
        rules = [dict(row) for row in cursor.fetchall()]

        # Categories & Institutions metadata
        cursor.execute("SELECT * FROM categories")
        categories = [dict(row) for row in cursor.fetchall()]
        for c in categories:
            if c.get("fields_definition"):
                try:
                    c["fields_definition"] = json.loads(c["fields_definition"])
                except:
                    pass

        cursor.execute("SELECT * FROM institutions ORDER BY name ASC")
        institutions = [dict(row) for row in cursor.fetchall()]

        cursor.execute("SELECT * FROM people ORDER BY name ASC")
        people = [dict(row) for row in cursor.fetchall()]

        cursor.execute("""
            SELECT tp.*, p.name as person_name
            FROM tax_profiles tp
            LEFT JOIN people p ON tp.person_id = p.id
            ORDER BY tp.financial_year DESC, p.name ASC
        """)
        tax_profiles = [dict(row) for row in cursor.fetchall()]

        cursor.execute("SELECT * FROM card_variants ORDER BY is_seeded DESC, name ASC")
        card_variants = [dict(row) for row in cursor.fetchall()]

        cursor.execute("SELECT * FROM knowledge_articles ORDER BY topic ASC, title ASC")
        knowledge_articles = [_shape_knowledge_row(row) for row in cursor.fetchall()]

        cursor.execute("SELECT provider FROM ai_settings WHERE id = 1")
        ai_row = cursor.fetchone()
        ai_provider = ai_row["provider"] if ai_row else "none"

        # Monthly Cash Flows logs
        cursor.execute("SELECT * FROM monthly_cash_flows ORDER BY month DESC")
        cash_flows = [dict(row) for row in cursor.fetchall()]

        # Dynamic Recommendations mapping
        recommendations = generate_recommendations(conn)

        # Summary aggregates
        total_liquid = sum(float(a.get("current_balance") if a.get("current_balance") is not None else 0.0) for a in accounts)
        total_assets_val = sum(float(a.get("current_value") if a.get("current_value") is not None else 0.0) for a in assets)
        total_debt = sum(float(l.get("outstanding") if l.get("outstanding") is not None else 0.0) for l in liabilities)
        
        total_assets = total_liquid + total_assets_val
        net_worth = total_assets - total_debt

        # Check historic change (compare with last month snapshot)
        cursor.execute("SELECT net_worth FROM snapshots ORDER BY month DESC LIMIT 2")
        historic_snapshots = cursor.fetchall()
        
        monthly_change = 0.0
        yearly_change = 0.0
        
        if len(historic_snapshots) >= 1:
            monthly_change = net_worth - historic_snapshots[0]["net_worth"]
        
        cursor.execute("SELECT net_worth FROM snapshots WHERE month <= ? ORDER BY month DESC LIMIT 1", 
                       (str(datetime.now().year - 1),))
        last_year_snap = cursor.fetchone()
        if last_year_snap:
            yearly_change = net_worth - last_year_snap["net_worth"]

        summary = {
            "net_worth": net_worth,
            "total_assets": total_assets,
            "total_debt": total_debt,
            "total_liquid": total_liquid,
            "total_investments": sum(a["current_value"] for a in assets if a["category_name"] not in ["Property", "Insurance"]),
            "total_properties": sum(a["current_value"] for a in assets if a["category_name"] == "Property"),
            "monthly_change": monthly_change,
            "yearly_change": yearly_change
        }

        return {
            "summary": summary,
            "accounts": accounts,
            "assets": assets,
            "liabilities": liabilities,
            "income": income,
            "expenses": expenses,
            "goals": goals,
            "rules": rules,
            "categories": categories,
            "institutions": institutions,
            "people": people,
            "tax_profiles": tax_profiles,
            "card_variants": card_variants,
            "knowledge_articles": knowledge_articles,
            "ai_provider": ai_provider,
            "recommendations": recommendations,
            "cash_flows": cash_flows
        }
    finally:
        conn.close()

def normalize_legacy_transaction_payload(payload: dict):
    """
    Map the Transactions Ledger's legacy payload shape (account_id + type +
    related_entity_type/related_entity_id) onto the double-entry columns
    (source_ref_type/source_ref_id/dest_ref_type/dest_ref_id). No-op when the
    payload already carries source_ref_type. Must run on BOTH create and update:
    update reverts the old row's balance impact then reapplies from the new
    payload, so a payload without refs would revert without reapplying.
    """
    if "source_ref_type" in payload:
        return

    acc_id = payload.get("account_id")
    ttype = payload.get("type")
    rel_id = payload.get("related_entity_id")

    source_type = 'None'
    source_id = None
    dest_type = 'None'
    dest_id = None
    new_type = 'Credit'

    if ttype == 'Credit':
        source_type = 'None'
        dest_type = 'Account'
        dest_id = acc_id
        new_type = 'Credit'
    elif ttype == 'Debit':
        source_type = 'Account'
        source_id = acc_id
        dest_type = 'None'
        new_type = 'Debit'
    elif ttype == 'Transfer':
        source_type = 'Account'
        source_id = acc_id
        dest_type = 'Account'
        dest_id = rel_id
        new_type = 'Transfer'
    elif ttype in ('Investment', 'Buy'):
        source_type = 'Account'
        source_id = acc_id
        dest_type = 'Asset'
        dest_id = rel_id
        new_type = 'Buy'
    elif ttype in ('Redemption', 'Sell'):
        source_type = 'Asset'
        source_id = rel_id
        dest_type = 'Account'
        dest_id = acc_id
        new_type = 'Sell'
    elif ttype in ('EMI', 'Prepayment'):
        source_type = 'Account'
        source_id = acc_id
        dest_type = 'Liability'
        dest_id = rel_id
        new_type = ttype
    elif ttype in ('Interest', 'Fee'):
        # A charge on a loan or card (interest for the month, a processing fee or a
        # penalty): it raises what's owed and moves no cash; the EMI pays it off.
        source_type = 'None'
        source_id = None
        dest_type = 'Liability'
        dest_id = rel_id
        new_type = ttype
    elif ttype == 'Spend':
        # A credit card charge: no cash account is touched at the moment of a
        # swipe (unlike EMI/Prepayment, which pay down the card from a real
        # Account) — it purely increases what's owed on the card.
        source_type = 'None'
        source_id = None
        dest_type = 'Liability'
        dest_id = rel_id
        new_type = 'Spend'

    payload["source_ref_type"] = source_type
    payload["source_ref_id"] = source_id
    payload["dest_ref_type"] = dest_type
    payload["dest_ref_id"] = dest_id
    payload["type"] = new_type

# Human-readable label for the one synthetic 'Spend' transaction per credit
# card representing "spend I haven't logged individually this cycle" (see
# put_credit_card_details and init_db()'s one-time migration of legacy
# current_statement_spend values). Display/lookup convenience only — the
# balance exemption itself is enforced via the real is_balance_exempt column
# (transactions.is_balance_exempt), NOT by matching this string. An earlier
# version tried exact-matching this free-text notes value instead, which was
# unsafe two ways: (1) editing the marker row's own Notes text could silently
# un-exempt it mid-edit, and (2) if any row's stored notes ever happened to
# equal this string, every future edit/delete of that row would incorrectly
# skip reverting its real balance impact, permanently. is_balance_exempt is a
# real column the generic entity CRUD route never writes unless the caller's
# payload explicitly includes it — and the Transactions Ledger UI never does
# for any transaction type — so a normal user edit structurally cannot change
# it, closing both failure modes at the source instead of requiring every
# call site to carefully reason about which of "old" vs "new" text to trust.
STATEMENT_SPEND_ADJUSTMENT_NOTE = "Manual statement spend adjustment"

def adjust_double_entry_balances(cursor, t_type: str, amount: float, source_type: str, source_id: int, dest_type: str, dest_id: int, operation: str = "add", is_balance_exempt=False):
    """
    Adjust balances of assets and liabilities based on double-entry transaction specs.
    (Account balances are derived dynamically from transaction history, so they are not stored/updated).
    operation: 'add' (apply transaction impact), 'revert' (undo transaction impact)
    """
    if not source_type and not dest_type:
        return

    factor = 1.0 if operation == "add" else -1.0

    # 1. Update Asset values
    if source_type == 'Asset' and source_id:
        cursor.execute("UPDATE assets SET current_value = current_value - ? WHERE id = ?", (amount * factor, source_id))
    if dest_type == 'Asset' and dest_id:
        cursor.execute("UPDATE assets SET current_value = current_value + ? WHERE id = ?", (amount * factor, dest_id))

    # 2. Update Liability outstanding by the direction of the money, like accounts and
    # assets (the type is only a label: EMI, Prepayment, Spend, Interest, Fee...):
    #   account/asset -> liability  pays it down        (EMI, prepayment, card bill)
    #   outside       -> liability  is a charge on it   (card spend, loan interest, fees)
    #   liability     -> anywhere   is borrowing more   (a disbursal or top-up)
    # Every entry that existed before this rule had one of the first two shapes with the
    # same effect (EMI/Prepayment came from an account, Spend from outside), so existing
    # balances are unchanged (test_ledger_direction.py). A row flagged is_balance_exempt
    # (the statement-spend-adjustment marker; see its constant's docstring above) never
    # moves a balance.
    if is_balance_exempt:
        return
    if dest_type == 'Liability' and dest_id:
        paid_from_own_money = source_type in ('Account', 'Asset')
        cursor.execute("UPDATE liabilities SET outstanding = outstanding + ? WHERE id = ?",
                       ((-amount if paid_from_own_money else amount) * factor, dest_id))
    if source_type == 'Liability' and source_id:
        cursor.execute("UPDATE liabilities SET outstanding = outstanding + ? WHERE id = ?", (amount * factor, source_id))

def compute_real_cycle_spend_total(cursor, liability_id: int, statement_date) -> dict:
    """
    Sums 'Spend' transactions against this card within its current billing
    cycle (see calc.calculate_statement_cycle) — computed fresh on every call,
    no cached/stored total. has_any_ever gates the manual-vs-derived fallback
    in the frontend: a card with a pre-existing manually-entered
    current_statement_spend and zero logged Spend transactions must keep
    showing its old value, not silently drop to 0.
    """
    cycle = calc.calculate_statement_cycle(statement_date, date.today())
    cursor.execute("""
        SELECT COALESCE(SUM(amount), 0.0) AS total FROM transactions
        WHERE type = 'Spend' AND dest_ref_type = 'Liability' AND dest_ref_id = ?
          AND transaction_date >= ? AND transaction_date <= ?
    """, (liability_id, cycle["start"], cycle["end"]))
    total = cursor.fetchone()["total"]
    cursor.execute(
        "SELECT 1 FROM transactions WHERE type = 'Spend' AND dest_ref_type = 'Liability' AND dest_ref_id = ? LIMIT 1",
        (liability_id,)
    )
    has_any_ever = cursor.fetchone() is not None
    # The one marked adjustment transaction's own amount, separate from the full
    # cycle total (which also includes any other real logged Spend rows) — this
    # is what the "Statement Spend Adjustment" form field should prefill with,
    # not the aggregate.
    cursor.execute("""
        SELECT amount FROM transactions
        WHERE dest_ref_type = 'Liability' AND dest_ref_id = ? AND type = 'Spend'
          AND is_balance_exempt = 1
    """, (liability_id,))
    adj_row = cursor.fetchone()
    adjustment_amount = float(adj_row["amount"]) if adj_row else 0.0
    return {
        "total": round(total, 2), "has_any_ever": has_any_ever,
        "cycle_start": cycle["start"], "cycle_end": cycle["end"],
        "adjustment_amount": adjustment_amount,
    }

def record_capital_gains_realization(cursor, asset_id: int, transaction_id: int, sale_value: float, sale_date: str):
    """
    Called when a 'Sell' transaction is recorded against a Stock/Mutual Fund/Property
    asset. FIFO-consumes holding_lots (Stock/Mutual Fund only — Property has no lots,
    it's a single cost basis) up to the asset's current quantity, logs the resulting
    LTCG/STCG split into capital_gains_realizations, and reduces/clears the consumed lots.
    Best-effort estimate only, not tax advice — see calculations.py.
    """
    cursor.execute("""
        SELECT a.owner_id, a.quantity, a.cost_price, a.purchase_date, c.name as category_name
        FROM assets a JOIN categories c ON a.category_id = c.id
        WHERE a.id = ?
    """, (asset_id,))
    asset = cursor.fetchone()
    if not asset or asset["category_name"] not in ("Stock", "Mutual Fund", "Property"):
        return

    financial_year = calc.get_current_financial_year()

    if asset["category_name"] == "Property":
        cost_basis = float(asset["cost_price"] or 0.0)
        gain = sale_value - cost_basis
        gain_type = calc.determine_gain_type(asset["purchase_date"], sale_date, "Property") if asset["purchase_date"] else "STCG"
        cursor.execute("""
            INSERT INTO capital_gains_realizations
            (asset_id, transaction_id, owner_id, asset_category, sale_date, financial_year,
             quantity_sold, sale_value, cost_basis, gain_amount, gain_type)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (asset_id, transaction_id, asset["owner_id"], "Property", sale_date, financial_year,
              1.0, sale_value, cost_basis, gain, gain_type))
        return

    cursor.execute("SELECT * FROM holding_lots WHERE asset_id = ? ORDER BY purchase_date", (asset_id,))
    lots = [dict(row) for row in cursor.fetchall()]
    total_qty = sum(l["quantity"] for l in lots)
    if total_qty <= 0:
        return
    sale_price_per_unit = sale_value / total_qty
    result = calc.calculate_capital_gains_for_lots(lots, total_qty, sale_price_per_unit, sale_date, asset["category_name"])

    for consumed in result["lots_consumed"]:
        cursor.execute("""
            INSERT INTO capital_gains_realizations
            (asset_id, transaction_id, owner_id, asset_category, sale_date, financial_year,
             quantity_sold, sale_value, cost_basis, gain_amount, gain_type)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (asset_id, transaction_id, asset["owner_id"], asset["category_name"], sale_date, financial_year,
              consumed["quantity_consumed"], consumed["sale_value"], consumed["cost_basis"],
              consumed["gain"], consumed["gain_type"]))

    # Fully sold: clear all lots. This route always sells the entire current holding
    # (no partial-sale/lot-selection UI exists yet — see Taxation domain non-goals).
    cursor.execute("DELETE FROM holding_lots WHERE asset_id = ?", (asset_id,))

# --- Generic CRUD REST API ---
@app.get("/api/entities/{entity_type}")
def get_entities(entity_type: str):
    allowed_tables = ["accounts", "assets", "liabilities", "income", "expenses", "goals", "transactions"]
    if entity_type not in allowed_tables:
        raise HTTPException(status_code=400, detail=f"Invalid entity type: {entity_type}")
        
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        if entity_type == "accounts":
            cursor.execute("""
                SELECT 
                    a.*,
                    (
                        SELECT COALESCE(SUM(t.amount), 0.0) 
                        FROM transactions t 
                        WHERE t.dest_ref_type = 'Account' AND t.dest_ref_id = a.id
                    ) - (
                        SELECT COALESCE(SUM(t.amount), 0.0) 
                        FROM transactions t 
                        WHERE t.source_ref_type = 'Account' AND t.source_ref_id = a.id
                    ) AS current_balance
                FROM accounts a
            """)
        elif entity_type == "goals":
            # Aliased derived_current_amount, not current_amount: g.* already carries the
            # stored column and sqlite3.Row resolves duplicate names to the leftmost, so a
            # same-named alias would be shadowed. Derived only when holdings exist —
            # manually-valued goals (no linked assets) keep their stored amount.
            cursor.execute("""
                SELECT
                    g.*,
                    CASE WHEN EXISTS (SELECT 1 FROM goal_holdings gh WHERE gh.goal_id = g.id)
                        THEN COALESCE(
                            (
                                SELECT SUM(a.current_value * gh.allocation_weight)
                                FROM assets a
                                JOIN goal_holdings gh ON a.id = gh.asset_id
                                WHERE gh.goal_id = g.id
                            ), 0.0
                        )
                        ELSE g.current_amount
                    END AS derived_current_amount
                FROM goals g
            """)
        else:
            cursor.execute(f"SELECT * FROM {entity_type}")
        rows = cursor.fetchall()

        result = []
        for row in rows:
            d = dict(row)
            if entity_type == "goals" and "derived_current_amount" in d:
                d["current_amount"] = d.pop("derived_current_amount")
            for k in ["metadata_json", "funding_sources_json", "snapshot_data"]:
                if k in d and d[k]:
                    try:
                        d[k] = json.loads(d[k])
                    except:
                        pass
            result.append(d)
        return result
    finally:
        conn.close()

@app.post("/api/entities/{entity_type}")
def create_entity(entity_type: str, payload: dict):
    allowed_tables = ["accounts", "assets", "liabilities", "income", "expenses", "goals", "transactions"]
    if entity_type not in allowed_tables:
        raise HTTPException(status_code=400, detail=f"Invalid entity type: {entity_type}")
        
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        initial_balance = 0.0
        if entity_type in ("accounts", "assets", "liabilities", "income", "expenses") and not payload.get("owner_id"):
            payload["owner_id"] = default_owner_id(cursor)
        if entity_type == "accounts":
            initial_balance = float(payload.pop("current_balance", 0.0))

        if entity_type == "transactions":
            normalize_legacy_transaction_payload(payload)
            # is_balance_exempt is never a caller-settable field — it's only ever
            # written internally by put_credit_card_details's adjustment marker.
            # Strip it from any incoming payload so a spoofed/spread payload can't
            # create a transaction that silently skips its balance impact.
            payload.pop("is_balance_exempt", None)

        # Preprocess custom fields and map expected dates for Goals
        if entity_type == "goals" and "target_amount" in payload and "current_amount" in payload and "monthly_contribution" in payload:
            comp = calc.calculate_goal_expected_completion(
                float(payload.get("target_amount", 0.0)),
                float(payload.get("current_amount", 0.0)),
                float(payload.get("monthly_contribution", 0.0))
            )
            payload["expected_completion"] = comp["expected_date"]

        standard_params = serialize_entity_payload(cursor, entity_type, payload)

        columns = list(standard_params.keys())
        placeholders = [f":{c}" for c in columns]
        query = f"INSERT INTO {entity_type} ({', '.join(columns)}) VALUES ({', '.join(placeholders)})"
        
        cursor.execute(query, standard_params)
        
        entity_id = cursor.lastrowid
        
        if entity_type == "accounts" and initial_balance > 0.0:
            cursor.execute("""
            INSERT INTO transactions (source_ref_type, source_ref_id, dest_ref_type, dest_ref_id, type, amount, transaction_date, notes)
            VALUES ('None', NULL, 'Account', ?, 'Credit', ?, ?, 'Initial balance registration')
            """, (entity_id, initial_balance, date.today().isoformat()))
            
        if entity_type == "transactions":
            adjust_double_entry_balances(
                cursor,
                standard_params.get("type"),
                float(standard_params.get("amount", 0.0)),
                standard_params.get("source_ref_type"),
                int(standard_params.get("source_ref_id")) if standard_params.get("source_ref_id") is not None else None,
                standard_params.get("dest_ref_type"),
                int(standard_params.get("dest_ref_id")) if standard_params.get("dest_ref_id") is not None else None,
                "add",
                bool(standard_params.get("is_balance_exempt"))
            )
            if standard_params.get("type") == "Sell" and standard_params.get("source_ref_type") == "Asset" and standard_params.get("source_ref_id"):
                record_capital_gains_realization(
                    cursor,
                    int(standard_params["source_ref_id"]),
                    entity_id,
                    float(standard_params.get("amount", 0.0)),
                    standard_params.get("transaction_date") or date.today().isoformat()
                )

        if entity_type == "assets":
            goal_mapping = payload.get("goal_mapping")
            if goal_mapping:
                cursor.execute("SELECT id FROM goals WHERE name = ?", (goal_mapping,))
                g_row = cursor.fetchone()
                if g_row:
                    cursor.execute("""
                    INSERT OR REPLACE INTO goal_holdings (goal_id, asset_id, allocation_weight)
                    VALUES (?, ?, 1.0)
                    """, (g_row["id"], entity_id))

        conn.commit()
        return {"status": "success", "id": entity_id}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/entities/{entity_type}/{id}")
def update_entity(entity_type: str, id: int, payload: dict):
    allowed_tables = ["accounts", "assets", "liabilities", "income", "expenses", "goals", "transactions"]
    if entity_type not in allowed_tables:
        raise HTTPException(status_code=400, detail=f"Invalid entity type: {entity_type}")
        
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        if entity_type == "accounts":
            # Accounts have no stored balance (it's derived from the transactions ledger),
            # so an edited opening balance maps to the account's 'Initial balance
            # registration' credit transaction: update it in place, create it if the
            # account was registered without one, or delete it when set back to 0.
            # current_balance is popped too: it's the derived read-only figure echoed
            # back by the form and must never be written anywhere.
            payload.pop("current_balance", None)
            opening_raw = payload.pop("opening_balance", None)
            if opening_raw is not None:
                opening_balance = float(opening_raw or 0.0)
                # Two markers identify an account's opening credit: create_entity writes
                # 'Initial balance registration', the MMex CSV importer writes
                # 'CSV Import Initial Balance'. Matching both here (and in /api/init's
                # opening_balance subquery) prevents an edit from inserting a duplicate
                # opening transaction alongside an import-created one.
                cursor.execute("""
                    SELECT id, amount FROM transactions
                    WHERE dest_ref_type = 'Account' AND dest_ref_id = ?
                      AND notes IN ('Initial balance registration', 'CSV Import Initial Balance')
                    ORDER BY id LIMIT 1
                """, (id,))
                initial_tx = cursor.fetchone()
                if initial_tx is None and opening_balance > 0.0:
                    cursor.execute("""
                    INSERT INTO transactions (source_ref_type, source_ref_id, dest_ref_type, dest_ref_id, type, amount, transaction_date, notes)
                    VALUES ('None', NULL, 'Account', ?, 'Credit', ?, ?, 'Initial balance registration')
                    """, (id, opening_balance, date.today().isoformat()))
                elif initial_tx is not None and opening_balance > 0.0 and float(initial_tx["amount"]) != opening_balance:
                    cursor.execute("UPDATE transactions SET amount = ? WHERE id = ?", (opening_balance, initial_tx["id"]))
                elif initial_tx is not None and opening_balance == 0.0:
                    cursor.execute("DELETE FROM transactions WHERE id = ?", (initial_tx["id"],))

        # Preprocess custom fields and map expected dates for Goals
        if entity_type == "goals" and "target_amount" in payload and "current_amount" in payload and "monthly_contribution" in payload:
            comp = calc.calculate_goal_expected_completion(
                float(payload.get("target_amount", 0.0)),
                float(payload.get("current_amount", 0.0)),
                float(payload.get("monthly_contribution", 0.0))
            )
            payload["expected_completion"] = comp["expected_date"]
            
        if entity_type == "transactions":
            # The Ledger edits with the legacy payload shape — normalize it to the
            # double-entry columns BEFORE the revert/reapply cycle, otherwise the old
            # impact is reverted but the reapply below sees no refs and does nothing.
            normalize_legacy_transaction_payload(payload)
            # is_balance_exempt is never a caller-settable field (see create_entity)
            # — strip it so a payload can't flip a row's exemption status directly,
            # which would corrupt the NEXT edit's revert/reapply (both of which
            # trust old.get("is_balance_exempt") as the row's true stored value).
            payload.pop("is_balance_exempt", None)
            # Revert old balances first
            cursor.execute("SELECT * FROM transactions WHERE id = ?", (id,))
            old_row = cursor.fetchone()
            if old_row:
                old = dict(old_row)
                adjust_double_entry_balances(
                    cursor,
                    old.get("type"),
                    float(old.get("amount", 0.0)),
                    old.get("source_ref_type"),
                    int(old.get("source_ref_id")) if old.get("source_ref_id") is not None else None,
                    old.get("dest_ref_type"),
                    int(old.get("dest_ref_id")) if old.get("dest_ref_id") is not None else None,
                    "revert",
                    bool(old.get("is_balance_exempt"))
                )

        standard_params = serialize_entity_payload(cursor, entity_type, payload)

        set_clauses = [f"{c} = :{c}" for c in standard_params.keys()]
        standard_params["id"] = id
        query = f"UPDATE {entity_type} SET {', '.join(set_clauses)} WHERE id = :id"

        cursor.execute(query, standard_params)

        if entity_type == "transactions":
            adjust_double_entry_balances(
                cursor,
                standard_params.get("type"),
                float(standard_params.get("amount", 0.0)),
                standard_params.get("source_ref_type"),
                int(standard_params.get("source_ref_id")) if standard_params.get("source_ref_id") is not None else None,
                standard_params.get("dest_ref_type"),
                int(standard_params.get("dest_ref_id")) if standard_params.get("dest_ref_id") is not None else None,
                "add",
                # is_balance_exempt from `old` (the pre-edit row), not
                # standard_params: the generic Ledger's payload never includes this
                # column at all (it's not an editable form field), so
                # standard_params wouldn't have it regardless — the UPDATE above
                # also never touches it for the same reason, meaning `old`'s value
                # is guaranteed to still be this row's true current value here.
                bool(old.get("is_balance_exempt")) if old_row else False
            )

        if entity_type == "assets":
            cursor.execute("DELETE FROM goal_holdings WHERE asset_id = ?", (id,))
            goal_mapping = payload.get("goal_mapping")
            if goal_mapping:
                cursor.execute("SELECT id FROM goals WHERE name = ?", (goal_mapping,))
                g_row = cursor.fetchone()
                if g_row:
                    cursor.execute("""
                    INSERT OR REPLACE INTO goal_holdings (goal_id, asset_id, allocation_weight)
                    VALUES (?, ?, 1.0)
                    """, (g_row["id"], id))

        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/entities/{entity_type}/{id}")
def delete_entity(entity_type: str, id: int):
    allowed_tables = ["accounts", "assets", "liabilities", "income", "expenses", "goals", "transactions"]
    if entity_type not in allowed_tables:
        raise HTTPException(status_code=400, detail=f"Invalid entity type: {entity_type}")
        
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        if entity_type == "transactions":
            # Revert old balances first
            cursor.execute("SELECT * FROM transactions WHERE id = ?", (id,))
            old_row = cursor.fetchone()
            if old_row:
                old = dict(old_row)
                adjust_double_entry_balances(
                    cursor,
                    old.get("type"),
                    float(old.get("amount", 0.0)),
                    old.get("source_ref_type"),
                    int(old.get("source_ref_id")) if old.get("source_ref_id") is not None else None,
                    old.get("dest_ref_type"),
                    int(old.get("dest_ref_id")) if old.get("dest_ref_id") is not None else None,
                    "revert",
                    bool(old.get("is_balance_exempt"))
                )

        if entity_type == "assets":
            cursor.execute("DELETE FROM goal_holdings WHERE asset_id = ?", (id,))

        doc_files_to_remove = []
        if entity_type in DOCUMENT_REF_TYPE_MAP:
            ref_type = DOCUMENT_REF_TYPE_MAP[entity_type]
            cursor.execute("SELECT stored_name FROM documents WHERE ref_type = ? AND ref_id = ?", (ref_type, id))
            doc_files_to_remove = [row["stored_name"] for row in cursor.fetchall()]
            cursor.execute("DELETE FROM documents WHERE ref_type = ? AND ref_id = ?", (ref_type, id))

        cursor.execute(f"DELETE FROM {entity_type} WHERE id = ?", (id,))
        conn.commit()
        for stored_name in doc_files_to_remove:
            file_path = os.path.join(get_documents_dir(), stored_name)
            if os.path.exists(file_path):
                os.remove(file_path)
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- Satellite Details Routers (Credit Card, Insurance, Holdings) ---
class CreditCardDetailsSchema(BaseModel):
    credit_limit: float = 0.0
    statement_balance: float = 0.0
    statement_date: Optional[int] = None
    due_date: Optional[int] = None
    annual_fee: float = 0.0
    reward_points: float = 0.0
    point_value: float = 0.25
    renewal_date: Optional[str] = None
    network: str = "Visa"
    lounge_visits_limit: float = 0.0
    lounge_visits_used: float = 0.0
    lounge_int_limit: float = 0.0
    lounge_int_used: float = 0.0
    milestone_spend_target: float = 0.0
    milestone_spend_current: float = 0.0
    milestone_reward_desc: Optional[str] = None
    remarks: Optional[str] = None
    base_points: float = 1.0
    base_spend: float = 100.0
    current_statement_spend: float = 0.0
    category_rules_json: str = "[]"
    milestone_rules_json: str = "[]"
    redemption_partners_json: str = "[]"

@app.get("/api/liabilities/{id}/credit-card-details")
def get_credit_card_details(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM credit_card_details WHERE liability_id = ?", (id,))
        row = cursor.fetchone()
        return dict(row) if row else None
    finally:
        conn.close()

class LoanPropertyLink(BaseModel):
    asset_id: Optional[int] = None


@app.put("/api/liabilities/{id}/property")
def set_loan_property(id: int, link: LoanPropertyLink):
    """Set or clear the property a loan is secured on (liabilities.linked_asset_id).

    A property can carry several loans (a home loan and a top-up, say), so the link
    lives on the loan. This touches only that column: the generic entity PUT rebuilds
    metadata_json from the payload, so a partial save through it would wipe the loan's
    other fields.
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        loan = cursor.execute("""SELECT c.name AS category FROM liabilities l LEFT JOIN categories c ON c.id = l.category_id
                                 WHERE l.id = ?""", (id,)).fetchone()
        if not loan:
            raise HTTPException(status_code=404, detail="Loan not found")
        if loan["category"] == "Credit Card":
            raise HTTPException(status_code=422, detail="A credit card isn't secured on a property")
        if link.asset_id is not None:
            asset = cursor.execute("""SELECT c.name AS category FROM assets a LEFT JOIN categories c ON c.id = a.category_id
                                      WHERE a.id = ?""", (link.asset_id,)).fetchone()
            if not asset:
                raise HTTPException(status_code=404, detail="Property not found")
            if asset["category"] != "Property":
                raise HTTPException(status_code=422, detail="A loan can only be linked to a Property")
        cursor.execute("UPDATE liabilities SET linked_asset_id = ? WHERE id = ?", (link.asset_id, id))
        conn.commit()
        return {"status": "success"}
    finally:
        conn.close()


@app.put("/api/liabilities/{id}/credit-card-details")
def put_credit_card_details(id: int, details: CreditCardDetailsSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT liability_id FROM credit_card_details WHERE liability_id = ?", (id,))
        exists = cursor.fetchone()
        payload = details.model_dump(exclude_unset=True) if exists else details.model_dump()

        # current_statement_spend is no longer written to the column (real spend is
        # now derived from the transactions ledger — see compute_real_cycle_spend_total
        # and get_init_data). It's instead a "spend I haven't logged individually this
        # cycle" adjustment, mapped onto one dated 'Spend' transaction flagged
        # is_balance_exempt=1 (same create/update/delete-when-zeroed pattern as
        # accounts' opening_balance). Editing it refreshes the transaction's date
        # to today, so it keeps counting toward whichever cycle is currently open.
        # Popped from `payload` unconditionally so the column is never written.
        #
        # This marker transaction NEVER affects liabilities.outstanding — enforced
        # centrally in adjust_double_entry_balances via the is_balance_exempt
        # column, not by sign-flipping logic here. Earlier versions of this code
        # tried making outstanding move in one direction here and the opposite in
        # db.py's one-time legacy-value migration (unsound: the two couldn't agree
        # on how much was "already applied" once a migration-created row was later
        # edited here), and later tried identifying the marker by exact-matching
        # its free-text notes (unsound: editing the Notes field in the generic
        # Ledger could silently change a row's exemption status). is_balance_exempt
        # is a real column the generic Ledger route never writes, so a normal edit
        # structurally cannot flip it — removing the inconsistency instead of
        # trying to track it.
        adjustment_provided = "current_statement_spend" in payload
        adjustment_raw = payload.pop("current_statement_spend", None)
        if adjustment_provided:
            new_amount = float(adjustment_raw or 0.0)
            cursor.execute("""
                SELECT id, amount FROM transactions
                WHERE dest_ref_type = 'Liability' AND dest_ref_id = ? AND type = 'Spend'
                  AND is_balance_exempt = 1
            """, (id,))
            existing_adj = cursor.fetchone()
            today_str = date.today().isoformat()
            if existing_adj is None and new_amount > 0.0:
                cursor.execute("""
                    INSERT INTO transactions (source_ref_type, source_ref_id, dest_ref_type, dest_ref_id, type, amount, transaction_date, notes, is_balance_exempt)
                    VALUES ('None', NULL, 'Liability', ?, 'Spend', ?, ?, ?, 1)
                """, (id, new_amount, today_str, STATEMENT_SPEND_ADJUSTMENT_NOTE))
            elif existing_adj is not None and new_amount > 0.0:
                # Only bump the date when the amount actually changed. The dedicated
                # Credit Cards modal resubmits this field on every save (even ones
                # only touching unrelated fields like lounge visits), and re-dating
                # on a no-op resubmit would resurrect an already-cycled-out
                # adjustment into the currently open cycle — reintroducing the
                # "baseline double-counts every cycle" bug this dated-row design
                # exists to avoid.
                if float(existing_adj["amount"]) != new_amount:
                    cursor.execute(
                        "UPDATE transactions SET amount = ?, transaction_date = ? WHERE id = ?",
                        (new_amount, today_str, existing_adj["id"])
                    )
            elif existing_adj is not None and new_amount == 0.0:
                cursor.execute("DELETE FROM transactions WHERE id = ?", (existing_adj["id"],))

            # The raw column is vestigial once this marker-transaction mechanism
            # exists — real spend is always read from the ledger (see
            # compute_real_cycle_spend_total), never from this column. But it must
            # still be explicitly zeroed on write (not just excluded from the SET
            # clause), otherwise a legacy non-zero value survives forever: the
            # frontend's has_spend_transactions ? real_cycle_spend_total :
            # current_statement_spend fallback would revert to the stale legacy
            # number the instant the marker transaction has none of its own, and
            # migration L's `WHERE current_statement_spend > 0` guard would treat
            # the row as still needing conversion on every future restart.
            payload["current_statement_spend"] = 0.0

        if exists:
            # Only overwrite fields the caller actually sent. The generic Liabilities
            # grid (EntityGrid.tsx) PUTs a partial payload (credit_limit/statement_balance/
            # statement_date/due_date/annual_fee/reward_points/point_value only) — a full
            # .model_dump() would fill in every omitted field from this schema's defaults and
            # silently wipe category_rules_json/milestone_rules_json/redemption_partners_json/
            # base_points/base_spend/etc. back to '[]'/0, destroying whatever was configured
            # via the dedicated Credit Cards screen.
            if payload:
                set_clauses = [f"{c} = :{c}" for c in payload.keys()]
                payload["id"] = id
                cursor.execute(f"UPDATE credit_card_details SET {', '.join(set_clauses)} WHERE liability_id = :id", payload)
        else:
            # New row: every column needs a real value, not just whatever the caller sent.
            payload["current_statement_spend"] = 0.0
            payload["liability_id"] = id
            columns = list(payload.keys())
            placeholders = [f":{c}" for c in columns]
            cursor.execute(
                f"INSERT INTO credit_card_details ({', '.join(columns)}) VALUES ({', '.join(placeholders)})",
                payload
            )
        cursor.execute("UPDATE liabilities SET liability_kind = 'revolving_credit', interest_rate = 0, emi = 0, tenure_months = 0 WHERE id = ?", (id,))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.get("/api/liabilities/{id}/spend-breakdown")
def get_spend_breakdown(id: int):
    """
    Per-category spend within the credit card's currently-open billing cycle,
    computed live from the transactions ledger — see compute_real_cycle_spend_total
    for the matching aggregate total surfaced in /api/init. Deliberately its own
    on-demand endpoint rather than baked into /api/init (same precedent as the
    Annual Review generator): only needed when a user actually opens a card's
    detail view, not on every app-wide refresh.
    """
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT statement_date FROM credit_card_details WHERE liability_id = ?", (id,))
        details_row = cursor.fetchone()
        statement_date = details_row["statement_date"] if details_row else None
        cycle = calc.calculate_statement_cycle(statement_date, date.today())
        cursor.execute("""
            SELECT c.name AS category_name, COALESCE(SUM(t.amount), 0.0) AS amount
            FROM transactions t
            LEFT JOIN categories c ON t.category_id = c.id
            WHERE t.type = 'Spend' AND t.dest_ref_type = 'Liability' AND t.dest_ref_id = ?
              AND t.transaction_date >= ? AND t.transaction_date <= ?
            GROUP BY c.name
            ORDER BY amount DESC
        """, (id, cycle["start"], cycle["end"]))
        by_category = [{"category_name": row["category_name"] or "Uncategorized", "amount": round(row["amount"], 2)} for row in cursor.fetchall()]
        return {
            "cycle_start": cycle["start"],
            "cycle_end": cycle["end"],
            "total": round(sum(r["amount"] for r in by_category), 2),
            "by_category": by_category,
        }
    finally:
        conn.close()

# --- Card Library (reference catalogue, distinct from owned cards above) ---
class CardVariantSchema(BaseModel):
    name: str
    issuer: str
    network: str = "Visa"
    card_tier: str = "Premium"
    joining_fee: float = 0.0
    annual_fee: float = 0.0
    fee_waiver_spend: float = 0.0
    base_points: float = 1.0
    base_spend: float = 100.0
    point_value: float = 0.25
    lounge_visits_domestic: int = 0
    lounge_visits_international: int = 0
    eligibility_notes: Optional[str] = None
    category_rules_json: str = "[]"
    milestone_rules_json: str = "[]"
    redemption_partners_json: str = "[]"
    pros_json: str = "[]"
    cons_json: str = "[]"
    source_notes: Optional[str] = None

@app.get("/api/card-variants")
def list_card_variants():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM card_variants ORDER BY is_seeded DESC, name ASC")
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/card-variants")
def create_card_variant(variant: CardVariantSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        payload = variant.model_dump()
        payload["is_seeded"] = 0
        columns = list(payload.keys())
        placeholders = [f":{c}" for c in columns]
        cursor.execute(
            f"INSERT INTO card_variants ({', '.join(columns)}) VALUES ({', '.join(placeholders)})",
            payload
        )
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/card-variants/{id}")
def update_card_variant(id: int, variant: CardVariantSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT id FROM card_variants WHERE id = ?", (id,))
        if not cursor.fetchone():
            raise HTTPException(status_code=404, detail="Card variant not found")
        payload = variant.model_dump()
        payload["id"] = id
        # Editing a starter-catalogue card makes it the user's own copy: is_seeded
        # drops to 0 so "Starter Catalogue" in the UI always means untouched shipped
        # reference data. (Seeding runs only on table creation, so the pristine row
        # is not resurrected either — see init_db in db.py.)
        payload["is_seeded"] = 0
        set_clauses = [f"{c} = :{c}" for c in payload.keys() if c != "id"]
        cursor.execute(
            f"UPDATE card_variants SET {', '.join(set_clauses)}, updated_at = datetime('now') WHERE id = :id",
            payload
        )
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/card-variants/{id}")
def delete_card_variant(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM card_variants WHERE id = ?", (id,))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- Knowledge Base (read-only reference content, no CRUD) ---
def _shape_knowledge_row(row) -> dict:
    d = dict(row)
    try:
        d["tags"] = json.loads(d.pop("tags_json") or "[]")
    except Exception:
        d["tags"] = []
    try:
        d["see_also"] = json.loads(d.pop("see_also_json") or "[]")
    except Exception:
        d["see_also"] = []
    return d

@app.get("/api/knowledge-articles")
def list_knowledge_articles(topic: Optional[str] = None, tag: Optional[str] = None):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM knowledge_articles ORDER BY topic ASC, title ASC")
        articles = [_shape_knowledge_row(row) for row in cursor.fetchall()]
        if topic:
            articles = [a for a in articles if a["topic"] == topic]
        if tag:
            articles = [a for a in articles if tag in a["tags"]]
        return articles
    finally:
        conn.close()

@app.get("/api/knowledge-articles/by-rule/{rule_key}")
def get_knowledge_article_by_rule(rule_key: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM knowledge_articles WHERE related_rule_key = ?", (rule_key,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="No article linked to this rule")
        return _shape_knowledge_row(row)
    except HTTPException:
        raise
    finally:
        conn.close()

@app.get("/api/knowledge-articles/{slug}")
def get_knowledge_article(slug: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM knowledge_articles WHERE slug = ?", (slug,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Article not found")
        return _shape_knowledge_row(row)
    except HTTPException:
        raise
    finally:
        conn.close()

class InsuranceDetailsSchema(BaseModel):
    policy_type: str = "Life"
    policy_number: Optional[str] = None
    premium: float = 0.0
    frequency: str = "Yearly"
    renewal_date: Optional[str] = None
    sum_assured: float = 0.0
    surrender_value: float = 0.0
    sum_insured: float = 0.0
    claims_made: int = 0
    claims_amount: float = 0.0
    idv: float = 0.0
    vehicle_registration: Optional[str] = None

@app.get("/api/assets/{id}/insurance-details")
def get_insurance_details(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM insurance_details WHERE asset_id = ?", (id,))
        row = cursor.fetchone()
        return dict(row) if row else None
    finally:
        conn.close()

@app.put("/api/assets/{id}/insurance-details")
def put_insurance_details(id: int, details: InsuranceDetailsSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM insurance_details WHERE asset_id = ?", (id,))
        exists = cursor.fetchone()
        payload = details.model_dump()
        if exists:
            set_clauses = [f"{c} = :{c}" for c in payload.keys()]
            payload["id"] = id
            cursor.execute(f"UPDATE insurance_details SET {', '.join(set_clauses)} WHERE asset_id = :id", payload)
        else:
            payload["asset_id"] = id
            columns = list(payload.keys())
            placeholders = [f":{c}" for c in columns]
            cursor.execute(
                f"INSERT INTO insurance_details ({', '.join(columns)}) VALUES ({', '.join(placeholders)})",
                payload
            )
        # An insurance policy's realizable value is its surrender value — keep the
        # asset's current_value in sync so net worth and allocation reflect it.
        cursor.execute("UPDATE assets SET current_value = ? WHERE id = ?", (details.surrender_value, id))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

def recompute_holding_aggregates(cursor, asset_id: int):
    cursor.execute("SELECT quantity, price_per_unit FROM holding_lots WHERE asset_id = ?", (asset_id,))
    lots = cursor.fetchall()
    total_qty = sum(lot["quantity"] for lot in lots)
    total_cost = sum(lot["quantity"] * lot["price_per_unit"] for lot in lots)
    avg_cost = (total_cost / total_qty) if total_qty > 0 else 0.0
    cursor.execute("UPDATE assets SET quantity = ?, cost_price = ? WHERE id = ?", (total_qty, avg_cost, asset_id))

class HoldingLotSchema(BaseModel):
    purchase_date: str
    quantity: float
    price_per_unit: float
    notes: Optional[str] = ""

@app.get("/api/assets/{id}/holdings")
def get_holdings(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM holding_lots WHERE asset_id = ? ORDER BY purchase_date", (id,))
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/assets/{id}/holdings")
def create_holding(id: int, lot: HoldingLotSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO holding_lots (asset_id, purchase_date, quantity, price_per_unit, notes) VALUES (?, ?, ?, ?, ?)",
            (id, lot.purchase_date, lot.quantity, lot.price_per_unit, lot.notes)
        )
        recompute_holding_aggregates(cursor, id)
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/holdings/{lot_id}")
def update_holding(lot_id: int, lot: HoldingLotSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM holding_lots WHERE id = ?", (lot_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Holding lot not found")
        asset_id = row["asset_id"]
        cursor.execute(
            "UPDATE holding_lots SET purchase_date = ?, quantity = ?, price_per_unit = ?, notes = ? WHERE id = ?",
            (lot.purchase_date, lot.quantity, lot.price_per_unit, lot.notes, lot_id)
        )
        recompute_holding_aggregates(cursor, asset_id)
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/holdings/{lot_id}")
def delete_holding(lot_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM holding_lots WHERE id = ?", (lot_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Holding lot not found")
        asset_id = row["asset_id"]
        cursor.execute("DELETE FROM holding_lots WHERE id = ?", (lot_id,))
        recompute_holding_aggregates(cursor, asset_id)
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

class RentalLedgerEntrySchema(BaseModel):
    month: str
    rent_amount: float = 0.0
    expenses_amount: float = 0.0
    notes: Optional[str] = ""

@app.get("/api/assets/{id}/rental-entries")
def get_rental_entries(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM property_rental_ledger WHERE asset_id = ? ORDER BY month", (id,))
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/assets/{id}/rental-entries")
def create_rental_entry(id: int, entry: RentalLedgerEntrySchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO property_rental_ledger (asset_id, month, rent_amount, expenses_amount, notes) VALUES (?, ?, ?, ?, ?)",
            (id, entry.month, entry.rent_amount, entry.expenses_amount, entry.notes)
        )
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/rental-entries/{entry_id}")
def update_rental_entry(entry_id: int, entry: RentalLedgerEntrySchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM property_rental_ledger WHERE id = ?", (entry_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Rental entry not found")
        cursor.execute(
            "UPDATE property_rental_ledger SET month = ?, rent_amount = ?, expenses_amount = ?, notes = ? WHERE id = ?",
            (entry.month, entry.rent_amount, entry.expenses_amount, entry.notes, entry_id)
        )
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/rental-entries/{entry_id}")
def delete_rental_entry(entry_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM property_rental_ledger WHERE id = ?", (entry_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Rental entry not found")
        cursor.execute("DELETE FROM property_rental_ledger WHERE id = ?", (entry_id,))
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

class InsuranceClaimSchema(BaseModel):
    claim_date: str
    claim_amount: float = 0.0
    status: Optional[str] = "Filed"
    description: Optional[str] = ""

@app.get("/api/assets/{id}/insurance-claims")
def get_insurance_claims(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM insurance_claims WHERE asset_id = ? ORDER BY claim_date", (id,))
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/assets/{id}/insurance-claims")
def create_insurance_claim(id: int, claim: InsuranceClaimSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO insurance_claims (asset_id, claim_date, claim_amount, status, description) VALUES (?, ?, ?, ?, ?)",
            (id, claim.claim_date, claim.claim_amount, claim.status, claim.description)
        )
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/insurance-claims/{claim_id}")
def update_insurance_claim(claim_id: int, claim: InsuranceClaimSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM insurance_claims WHERE id = ?", (claim_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Insurance claim not found")
        cursor.execute(
            "UPDATE insurance_claims SET claim_date = ?, claim_amount = ?, status = ?, description = ? WHERE id = ?",
            (claim.claim_date, claim.claim_amount, claim.status, claim.description, claim_id)
        )
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/insurance-claims/{claim_id}")
def delete_insurance_claim(claim_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM insurance_claims WHERE id = ?", (claim_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Insurance claim not found")
        cursor.execute("DELETE FROM insurance_claims WHERE id = ?", (claim_id,))
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

class InsurancePremiumPaymentSchema(BaseModel):
    payment_date: str
    amount: float = 0.0
    notes: Optional[str] = ""

@app.get("/api/assets/{id}/premium-payments")
def get_premium_payments(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM insurance_premium_payments WHERE asset_id = ? ORDER BY payment_date", (id,))
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/assets/{id}/premium-payments")
def create_premium_payment(id: int, payment: InsurancePremiumPaymentSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO insurance_premium_payments (asset_id, payment_date, amount, notes) VALUES (?, ?, ?, ?)",
            (id, payment.payment_date, payment.amount, payment.notes)
        )
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/premium-payments/{payment_id}")
def update_premium_payment(payment_id: int, payment: InsurancePremiumPaymentSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM insurance_premium_payments WHERE id = ?", (payment_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Premium payment not found")
        cursor.execute(
            "UPDATE insurance_premium_payments SET payment_date = ?, amount = ?, notes = ? WHERE id = ?",
            (payment.payment_date, payment.amount, payment.notes, payment_id)
        )
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/premium-payments/{payment_id}")
def delete_premium_payment(payment_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT asset_id FROM insurance_premium_payments WHERE id = ?", (payment_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Premium payment not found")
        cursor.execute("DELETE FROM insurance_premium_payments WHERE id = ?", (payment_id,))
        conn.commit()
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- Documents Router (generic file attachments on any entity) ---
DOCUMENT_REF_TYPE_MAP = {"accounts": "Account", "assets": "Asset", "liabilities": "Liability", "goals": "Goal"}

def get_documents_dir():
    documents_dir = os.path.join(get_db_dir(), "documents")
    os.makedirs(documents_dir, exist_ok=True)
    return documents_dir

@app.get("/api/entities/{entity_type}/{entity_id}/documents")
def list_documents(entity_type: str, entity_id: int):
    if entity_type not in DOCUMENT_REF_TYPE_MAP:
        raise HTTPException(status_code=400, detail=f"Invalid entity type: {entity_type}")
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "SELECT * FROM documents WHERE ref_type = ? AND ref_id = ? ORDER BY uploaded_at DESC",
            (DOCUMENT_REF_TYPE_MAP[entity_type], entity_id)
        )
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/entities/{entity_type}/{entity_id}/documents")
async def upload_document(entity_type: str, entity_id: int, file: UploadFile = File(...), description: str = Form("")):
    if entity_type not in DOCUMENT_REF_TYPE_MAP:
        raise HTTPException(status_code=400, detail=f"Invalid entity type: {entity_type}")
    contents = await file.read()
    stored_name = f"{uuid.uuid4().hex}{os.path.splitext(file.filename or '')[1]}"
    dest_path = os.path.join(get_documents_dir(), stored_name)
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        with open(dest_path, "wb") as f:
            f.write(contents)
        cursor.execute(
            """INSERT INTO documents (ref_type, ref_id, file_name, stored_name, mime_type, size_bytes, description)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (DOCUMENT_REF_TYPE_MAP[entity_type], entity_id, file.filename, stored_name, file.content_type, len(contents), description)
        )
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        if os.path.exists(dest_path):
            os.remove(dest_path)
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.get("/api/documents/{document_id}/download")
def download_document(document_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM documents WHERE id = ?", (document_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Document not found")
        doc = dict(row)
        file_path = os.path.join(get_documents_dir(), doc["stored_name"])
        if not os.path.exists(file_path):
            raise HTTPException(status_code=404, detail="Document file missing on disk")
        return FileResponse(file_path, filename=doc["file_name"], media_type=doc["mime_type"] or "application/octet-stream")
    finally:
        conn.close()

@app.delete("/api/documents/{document_id}")
def delete_document(document_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT stored_name FROM documents WHERE id = ?", (document_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Document not found")
        file_path = os.path.join(get_documents_dir(), row["stored_name"])
        cursor.execute("DELETE FROM documents WHERE id = ?", (document_id,))
        conn.commit()
        if os.path.exists(file_path):
            os.remove(file_path)
        return {"status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.get("/api/goals/{id}/funding-constraints")
def get_goal_funding_constraints(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            SELECT gfc.excluded_category_id, c.name as category_name
            FROM goal_funding_constraints gfc
            JOIN categories c ON c.id = gfc.excluded_category_id
            WHERE gfc.goal_id = ?
        """, (id,))
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

class GoalFundingConstraintSchema(BaseModel):
    excluded_category_id: int

@app.post("/api/goals/{id}/funding-constraints")
def create_goal_funding_constraint(id: int, constraint: GoalFundingConstraintSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT OR IGNORE INTO goal_funding_constraints (goal_id, excluded_category_id) VALUES (?, ?)",
            (id, constraint.excluded_category_id)
        )
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/goals/{id}/funding-constraints/{category_id}")
def delete_goal_funding_constraint(id: int, category_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "DELETE FROM goal_funding_constraints WHERE goal_id = ? AND excluded_category_id = ?",
            (id, category_id)
        )
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- Taxation Router ---
# Decision-support estimate only, NOT tax advice. All tax math lives in backend/tax/
# (engine.py). Rule sets are either shipped read-only files (backend/tax/rules/*.json)
# or user-edited copies stored in tax_rule_sets (backend/tax/rule_store.py); no slab,
# limit or rate is read from financial_rules or hardcoded here.

def _tax_store_call(fn, *args, **kwargs):
    """Runs a rule/scenario store call on a fresh connection, mapping its errors to HTTP."""
    conn = get_db_connection()
    try:
        return fn(conn, *args, **kwargs)
    except tax_store.RuleStoreError as e:
        raise HTTPException(status_code=e.status_code, detail=str(e))
    except tax_engine.TaxInputError as e:
        raise HTTPException(status_code=422, detail=str(e))
    finally:
        conn.close()

def _default_financial_year() -> str:
    return tax_rules.latest_financial_year(calc.get_current_financial_year())

class TaxComputeRequest(BaseModel):
    financial_year: Optional[str] = None
    rule_set_id: Optional[str] = None
    rule_set_version: Optional[int] = None
    # Shape is defined once, by backend/tax/engine.py's DEFAULT_INPUTS, and validated
    # by engine.normalize_inputs — not duplicated as a Pydantic schema here.
    inputs: dict = {}

def _compute_tax(conn, req: TaxComputeRequest) -> dict:
    fy = req.financial_year or (None if req.rule_set_id else _default_financial_year())
    rules, meta = tax_store.resolve(conn, fy, req.rule_set_id, req.rule_set_version)
    result = tax_engine.compute(req.inputs, rules, computed_at=datetime.now().isoformat(timespec="seconds"))
    result["rule_set"].update({"name": meta["name"], "source": meta["source"], "content_hash": meta["content_hash"]})
    return result

@app.post("/api/tax/compute")
def compute_tax(req: TaxComputeRequest):
    return _tax_store_call(_compute_tax, req)

def _resolve_for(conn, financial_year: Optional[str], rule_set_id: Optional[str], rule_set_version: Optional[int]):
    fy = financial_year or (None if rule_set_id else _default_financial_year())
    rules, meta = tax_store.resolve(conn, fy, rule_set_id, rule_set_version)
    return rules, meta

def _now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")

class TaxNpsScenariosRequest(TaxComputeRequest):
    levels_pct: Optional[List[float]] = None
    custom_pct: Optional[float] = None
    custom_monthly_amount: Optional[float] = None
    funding: str = "restructure"
    restructure_from: str = "special_allowance"
    salary_base_override: Optional[float] = None
    offered_pcts: Optional[List[float]] = None  # null = employer policy not confirmed
    regime: str = "best"

@app.post("/api/tax/nps-scenarios")
def tax_nps_scenarios(req: TaxNpsScenariosRequest):
    def run(conn):
        rules, meta = _resolve_for(conn, req.financial_year, req.rule_set_id, req.rule_set_version)
        out = tax_nps.nps_scenarios(req.inputs, rules, req.levels_pct, req.custom_pct, req.custom_monthly_amount,
                                    req.funding, req.restructure_from, req.salary_base_override, req.offered_pcts,
                                    req.regime, _now_iso())
        out["rule_set"].update({"name": meta["name"], "source": meta["source"]})
        return out
    return _tax_store_call(run)

class TaxPayslipRequest(TaxComputeRequest):
    optimization: Optional[dict] = None  # {employer_nps_pct | employer_nps_monthly, funding, restructure_from, salary_base_override}
    regime_before: str = "best"
    regime_after: str = "best"
    actual_monthly_tds_before: Optional[float] = None
    actual_monthly_tds_after: Optional[float] = None
    month_index: Optional[int] = None  # month-by-month entry: show this month (0 = April) instead of the average month

@app.post("/api/tax/payslip")
def tax_payslip(req: TaxPayslipRequest):
    def run(conn):
        rules, meta = _resolve_for(conn, req.financial_year, req.rule_set_id, req.rule_set_version)
        out = tax_payroll.payslip(req.inputs, rules, req.optimization, req.regime_before, req.regime_after,
                                  req.actual_monthly_tds_before, req.actual_monthly_tds_after, _now_iso(), req.month_index)
        out["rule_set"].update({"name": meta["name"], "source": meta["source"]})
        return out
    return _tax_store_call(run)

class TaxEpfSuggestionRequest(BaseModel):
    financial_year: str
    wages: List[float]  # EPF wages (basic + DA) for each month, April to March
    basis: str = "ceiling"  # "ceiling": statutory, wages capped at the dated ceiling; "full": on full wages

@app.post("/api/tax/epf-suggestion")
def tax_epf_suggestion(req: TaxEpfSuggestionRequest):
    try:
        return tax_epf.suggest_epf(req.financial_year, req.wages, req.basis)
    except tax_engine.TaxInputError as e:
        raise HTTPException(status_code=422, detail=str(e))

@app.get("/api/tax/prefill")
def tax_prefill(financial_year: Optional[str] = None, person_id: Optional[int] = None):
    """Proposed planner figures from Artha's own records; the UI applies only what the user ticks."""
    fy = financial_year or _default_financial_year()
    # The year's rules carry the cap on interest for home loans spent on repairs.
    return _tax_store_call(lambda conn: tax_prefill_mod.prefill(conn, fy, person_id, rules=_resolve_for(conn, fy, None, None)[0]))

class TaxAdvanceRequest(TaxComputeRequest):
    expected_tds: Optional[float] = None  # default: the tax on salary alone (what an employer deducts)
    regime: Optional[str] = None  # default: the profile's TDS regime, else the cheaper one

@app.post("/api/tax/advance-tax")
def tax_advance(req: TaxAdvanceRequest):
    def run(conn):
        rules, _meta = _resolve_for(conn, req.financial_year, req.rule_set_id, req.rule_set_version)
        return tax_advance_mod.advance_tax_schedule(req.inputs, rules, expected_tds=req.expected_tds, regime=req.regime)
    return _tax_store_call(run)

@app.get("/api/tax/harvesting")
def tax_harvesting(person_id: int, financial_year: Optional[str] = None):
    """Gain- and loss-harvesting suggestions from this person's equity purchase lots (read-only)."""
    def run(conn):
        rules, _meta = _resolve_for(conn, financial_year, None, None)
        return tax_harvesting_mod.harvest_plan(conn, person_id, rules["financial_year"], rules)
    return _tax_store_call(run)

PLANNER_SLOT_PATTERN = re.compile(r"^p(\d+)-(\d{4}-\d{2})$")

def _tax_calendar_events(today, horizon) -> list:
    """Advance-tax instalments and the return due date from each person's saved Tax Planner
    plans (tax_planner_drafts slots p<person>-<FY>), between today and the horizon."""
    events = []
    conn = get_db_connection()
    try:
        people = {r["id"]: r["name"] for r in conn.execute("SELECT id, name FROM people")}
        for row in conn.execute("SELECT slot, content_json FROM tax_planner_drafts").fetchall():
            m = PLANNER_SLOT_PATTERN.match(row["slot"])
            if not m:
                continue
            person, fy = people.get(int(m.group(1)), "Household"), m.group(2)
            try:
                content = json.loads(row["content_json"])
                rules, _meta = tax_store.resolve(conn, fy, None, None)
                sched = tax_advance_mod.advance_tax_schedule(content.get("inputs") or {}, rules, today=today)
            except Exception as e:  # a plan that no longer computes shouldn't break the calendar
                print(f"[Calendar] Skipped tax plan {row['slot']}: {e}")
                continue
            for inst in sched.get("instalments", []):
                due = date.fromisoformat(inst["due_date"])
                if inst["status"] == "upcoming" and inst["still_to_pay"] > 0 and today <= due <= horizon:
                    events.append({
                        "title": f"Advance tax ({person}): {inst['cumulative_pct']}% of the year's tax by this date",
                        "amount": inst["still_to_pay"],
                        "type": "Tax",
                        "date": due.isoformat(),
                        "details": f"{person} · FY {fy}, {sched['regime']} regime · from the Tax Planner",
                    })
            filing = sched.get("return_due_date")
            if filing and today <= date.fromisoformat(filing) <= horizon:
                # When advance tax applies, the instalments above already carry the unpaid tax;
                # only below the threshold (or for exempt seniors) is it paid at filing.
                unpaid = max(sched.get("advance_tax_liability", 0.0) - sched.get("paid_total", 0.0), 0.0) if sched.get("available") else 0.0
                at_filing = 0.0 if sched.get("required") else unpaid
                events.append({
                    "title": f"Income tax return due ({person}, FY {fy})",
                    "amount": at_filing,
                    "type": "Tax",
                    "date": filing,
                    "details": (f"{person} · self-assessment tax to pay before filing" if at_filing
                                else f"{person} · tax not covered by TDS is in the advance-tax instalments" if sched.get("required")
                                else f"{person} · from the Tax Planner"),
                })
    finally:
        conn.close()
    return events

PAYSLIP_PDF_MAX_FILES = 24
PAYSLIP_PDF_MAX_BYTES = 10 * 1024 * 1024

@app.post("/api/tax/payslip-pdf")
async def tax_payslip_pdf(files: List[UploadFile] = File(...), password: Optional[str] = Form(None)):
    """Reads payslip PDFs in memory (never stored) into a pay month and label/amount lines;
    the planner maps them onto the month-by-month grid. One result per file, in order."""
    if len(files) > PAYSLIP_PDF_MAX_FILES:
        raise HTTPException(status_code=422, detail=f"Up to {PAYSLIP_PDF_MAX_FILES} PDFs at a time.")
    results = []
    for f in files:
        data = await f.read(PAYSLIP_PDF_MAX_BYTES + 1)
        entry = {"file": f.filename}
        if len(data) > PAYSLIP_PDF_MAX_BYTES:
            results.append({**entry, "error": "too_large", "message": "Larger than 10 MB; this doesn't look like a payslip."})
            continue
        try:
            results.append({**entry, **tax_payslip_pdf_mod.read_payslip(data, password or None)})
        except tax_payslip_pdf_mod.PayslipPdfError as e:
            results.append({**entry, "error": e.code, "message": str(e)})
    return {"results": results}

PLANNER_DRAFT_SLOT = re.compile(r"^[a-z0-9_-]{1,40}$")

@app.get("/api/tax/planner-draft/{slot}")
def get_tax_planner_draft(slot: str):
    if not PLANNER_DRAFT_SLOT.match(slot):
        raise HTTPException(status_code=422, detail="Invalid draft slot")
    def fetch(conn):
        row = conn.execute("SELECT content_json, updated_at FROM tax_planner_drafts WHERE slot = ?", (slot,)).fetchone()
        return {"slot": slot, "content": json.loads(row["content_json"]) if row else None,
                "updated_at": row["updated_at"] if row else None}
    return _tax_store_call(fetch)

@app.put("/api/tax/planner-draft/{slot}")
def put_tax_planner_draft(slot: str, content: dict):
    if not PLANNER_DRAFT_SLOT.match(slot):
        raise HTTPException(status_code=422, detail="Invalid draft slot")
    def save(conn):
        conn.execute("""INSERT INTO tax_planner_drafts (slot, content_json, updated_at) VALUES (?, ?, datetime('now'))
                        ON CONFLICT(slot) DO UPDATE SET content_json = excluded.content_json, updated_at = excluded.updated_at""",
                     (slot, json.dumps(content)))
        conn.commit()
        return {"slot": slot, "saved": True}
    return _tax_store_call(save)

@app.delete("/api/tax/planner-draft/{slot}")
def delete_tax_planner_draft(slot: str):
    if not PLANNER_DRAFT_SLOT.match(slot):
        raise HTTPException(status_code=422, detail="Invalid draft slot")
    def remove(conn):
        conn.execute("DELETE FROM tax_planner_drafts WHERE slot = ?", (slot,))
        conn.commit()
        return {"slot": slot, "deleted": True}
    return _tax_store_call(remove)

class TaxSensitivityRequest(TaxComputeRequest):
    bonus_month: int = 11  # month-by-month mode: where an extra bonus lands (0 = April ... 11 = March)

@app.post("/api/tax/sensitivity")
def tax_sensitivity(req: TaxSensitivityRequest):
    def run(conn):
        if not 0 <= req.bonus_month <= 11:
            raise tax_engine.TaxInputError("bonus_month must be 0 (April) to 11 (March)")
        rules, _meta = _resolve_for(conn, req.financial_year, req.rule_set_id, req.rule_set_version)
        return tax_sensitivity_mod.sensitivity(req.inputs, rules, req.bonus_month)
    return _tax_store_call(run)

class TaxNpsTargetRequest(TaxComputeRequest):
    target_saving: float
    funding: str = "restructure"
    restructure_from: str = "special_allowance"
    regime: str = "best"
    salary_base_override: Optional[float] = None

@app.post("/api/tax/nps-target")
def tax_nps_target(req: TaxNpsTargetRequest):
    def run(conn):
        rules, _meta = _resolve_for(conn, req.financial_year, req.rule_set_id, req.rule_set_version)
        return tax_sensitivity_mod.nps_for_target_saving(req.inputs, rules, req.target_saving, req.funding,
                                                         req.restructure_from, req.regime, req.salary_base_override)
    return _tax_store_call(run)

class TaxWaterfallRequest(TaxComputeRequest):
    regime: str = "new"

@app.post("/api/tax/waterfall")
def tax_waterfall_route(req: TaxWaterfallRequest):
    def run(conn):
        rules, _meta = _resolve_for(conn, req.financial_year, req.rule_set_id, req.rule_set_version)
        return tax_waterfall.tax_waterfall(req.inputs, rules, req.regime)
    return _tax_store_call(run)

@app.get("/api/tax/input-metadata")
def get_tax_input_metadata():
    return {**tax_rules.load_input_metadata(), "default_financial_year": _default_financial_year()}

@app.get("/api/tax/rule-sets")
def get_tax_rule_sets():
    return _tax_store_call(tax_store.list_rule_sets)

@app.get("/api/tax/rule-sets/{rule_set_id}")
def get_tax_rule_set(rule_set_id: str, version: Optional[int] = None):
    def fetch(conn):
        content, meta = tax_store.get_rule_set(conn, rule_set_id, version)
        return {"meta": meta, "content": content}
    return _tax_store_call(fetch)

@app.get("/api/tax/rule-sets/{rule_set_id}/versions")
def get_tax_rule_set_versions(rule_set_id: str):
    return _tax_store_call(tax_store.list_versions, rule_set_id)

@app.get("/api/tax/rule-sets/{rule_set_id}/diff")
def diff_tax_rule_set(rule_set_id: str, version: Optional[int] = None, against_id: Optional[str] = None,
                      against_version: Optional[int] = None):
    """Changes from `against` (default: this set's base for v1 or previous version) to this set."""
    def run(conn):
        content, meta = tax_store.get_rule_set(conn, rule_set_id, version)
        if against_id:
            base, base_meta = tax_store.get_rule_set(conn, against_id, against_version)
        elif meta["source"] == "user" and meta["version"] > 1:
            base, base_meta = tax_store.get_rule_set(conn, rule_set_id, meta["version"] - 1)
        elif meta["based_on_rule_set_id"]:
            base, base_meta = tax_store.get_rule_set(conn, meta["based_on_rule_set_id"], meta["based_on_version"])
        else:
            return {"against": None, "changes": []}
        return {"against": base_meta, "changes": tax_store.diff_rule_sets(base, content)}
    return _tax_store_call(run)

class TaxRuleSetCreateRequest(BaseModel):
    name: str
    based_on_rule_set_id: Optional[str] = None
    based_on_version: Optional[int] = None
    financial_year: Optional[str] = None
    content: Optional[dict] = None
    note: Optional[str] = None

@app.post("/api/tax/rule-sets")
def create_tax_rule_set(req: TaxRuleSetCreateRequest):
    return _tax_store_call(tax_store.create_rule_set, req.name, req.based_on_rule_set_id, req.based_on_version,
                           req.content, req.financial_year, req.note)

class TaxRuleSetUpdateRequest(BaseModel):
    content: dict
    name: Optional[str] = None
    note: Optional[str] = None

@app.put("/api/tax/rule-sets/{rule_set_id}")
def update_tax_rule_set(rule_set_id: str, req: TaxRuleSetUpdateRequest):
    return _tax_store_call(tax_store.update_rule_set, rule_set_id, req.content, req.note, req.name)

@app.post("/api/tax/rule-sets/{rule_set_id}/archive")
def archive_tax_rule_set(rule_set_id: str):
    return _tax_store_call(tax_store.set_status, rule_set_id, "archived")

@app.post("/api/tax/rule-sets/{rule_set_id}/restore")
def restore_tax_rule_set(rule_set_id: str):
    return _tax_store_call(tax_store.set_status, rule_set_id, "active")

@app.delete("/api/tax/rule-sets/{rule_set_id}")
def delete_tax_rule_set(rule_set_id: str):
    _tax_store_call(tax_store.delete_rule_set, rule_set_id)
    return {"status": "success"}

class TaxRuleSetSelectionRequest(BaseModel):
    rule_set_id: Optional[str] = None  # null = use the shipped rule set for the year

@app.put("/api/tax/rule-set-selection/{financial_year}")
def select_tax_rule_set(financial_year: str, req: TaxRuleSetSelectionRequest):
    _tax_store_call(tax_store.set_selection, financial_year, req.rule_set_id)
    return {"status": "success"}

class TaxScenarioRequest(BaseModel):
    name: str
    inputs: dict = {}
    financial_year: Optional[str] = None
    rule_set_id: Optional[str] = None
    rule_set_version: Optional[int] = None
    profile_id: Optional[int] = None
    notes: Optional[str] = None

@app.get("/api/tax/scenarios")
def get_tax_scenarios():
    return _tax_store_call(tax_scenarios.list_scenarios)

@app.get("/api/tax/scenarios/{scenario_id}")
def get_tax_scenario(scenario_id: int):
    return _tax_store_call(tax_scenarios.get, scenario_id)

@app.post("/api/tax/scenarios")
def save_tax_scenario(req: TaxScenarioRequest):
    fy = req.financial_year or (None if req.rule_set_id else _default_financial_year())
    return _tax_store_call(tax_scenarios.save, req.name, req.inputs, fy, req.rule_set_id,
                           req.rule_set_version, req.profile_id, req.notes)

@app.post("/api/tax/scenarios/{scenario_id}/recalculate")
def recalculate_tax_scenario(scenario_id: int):
    return _tax_store_call(tax_scenarios.recalculate, scenario_id)

@app.delete("/api/tax/scenarios/{scenario_id}")
def delete_tax_scenario(scenario_id: int):
    _tax_store_call(tax_scenarios.delete, scenario_id)
    return {"status": "success"}

class TaxProfileSchema(BaseModel):
    person_id: int
    financial_year: str
    regime: str = "New"
    age_band: str = "Below60"
    hra_rent_paid: float = 0.0
    hra_basic_salary: float = 0.0
    hra_received: float = 0.0
    is_metro_city: bool = False
    home_loan_interest_self_occupied: float = 0.0
    other_income_reported: float = 0.0
    tds_deducted: float = 0.0
    advance_tax_paid: float = 0.0
    notes: Optional[str] = None

@app.get("/api/tax-profiles")
def get_tax_profiles(person_id: Optional[int] = None, financial_year: Optional[str] = None):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        query = """
            SELECT tp.*, p.name as person_name
            FROM tax_profiles tp
            LEFT JOIN people p ON tp.person_id = p.id
            WHERE 1=1
        """
        params = []
        if person_id is not None:
            query += " AND tp.person_id = ?"
            params.append(person_id)
        if financial_year is not None:
            query += " AND tp.financial_year = ?"
            params.append(financial_year)
        query += " ORDER BY tp.financial_year DESC, p.name ASC"
        cursor.execute(query, params)
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/tax-profiles")
def create_tax_profile(profile: TaxProfileSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            INSERT INTO tax_profiles
            (person_id, financial_year, regime, age_band, hra_rent_paid, hra_basic_salary, hra_received,
             is_metro_city, home_loan_interest_self_occupied, other_income_reported, tds_deducted,
             advance_tax_paid, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            profile.person_id, profile.financial_year, profile.regime, profile.age_band,
            profile.hra_rent_paid, profile.hra_basic_salary, profile.hra_received,
            int(profile.is_metro_city), profile.home_loan_interest_self_occupied,
            profile.other_income_reported, profile.tds_deducted, profile.advance_tax_paid, profile.notes
        ))
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/tax-profiles/{id}")
def update_tax_profile(id: int, profile: TaxProfileSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            UPDATE tax_profiles SET
                person_id = ?, financial_year = ?, regime = ?, age_band = ?, hra_rent_paid = ?,
                hra_basic_salary = ?, hra_received = ?, is_metro_city = ?,
                home_loan_interest_self_occupied = ?, other_income_reported = ?, tds_deducted = ?,
                advance_tax_paid = ?, notes = ?
            WHERE id = ?
        """, (
            profile.person_id, profile.financial_year, profile.regime, profile.age_band,
            profile.hra_rent_paid, profile.hra_basic_salary, profile.hra_received,
            int(profile.is_metro_city), profile.home_loan_interest_self_occupied,
            profile.other_income_reported, profile.tds_deducted, profile.advance_tax_paid, profile.notes, id
        ))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/tax-profiles/{id}")
def delete_tax_profile(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM tax_profiles WHERE id = ?", (id,))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

class TaxDeductionSchema(BaseModel):
    section: str
    label: str
    amount: float = 0.0
    notes: Optional[str] = None

@app.get("/api/tax-profiles/{id}/deductions")
def get_tax_deductions(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM tax_deductions WHERE tax_profile_id = ? ORDER BY section, id", (id,))
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/tax-profiles/{id}/deductions")
def create_tax_deduction(id: int, deduction: TaxDeductionSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO tax_deductions (tax_profile_id, section, label, amount, notes) VALUES (?, ?, ?, ?, ?)",
            (id, deduction.section, deduction.label, deduction.amount, deduction.notes)
        )
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/tax-deductions/{deduction_id}")
def update_tax_deduction(deduction_id: int, deduction: TaxDeductionSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "UPDATE tax_deductions SET section = ?, label = ?, amount = ?, notes = ? WHERE id = ?",
            (deduction.section, deduction.label, deduction.amount, deduction.notes, deduction_id)
        )
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/tax-deductions/{deduction_id}")
def delete_tax_deduction(deduction_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM tax_deductions WHERE id = ?", (deduction_id,))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.get("/api/tax-profiles/{id}/capital-gains-realizations")
def get_capital_gains_realizations(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT person_id, financial_year FROM tax_profiles WHERE id = ?", (id,))
        profile = cursor.fetchone()
        if not profile:
            raise HTTPException(status_code=404, detail="Tax profile not found")
        cursor.execute("""
            SELECT cgr.*, a.name as asset_name
            FROM capital_gains_realizations cgr
            LEFT JOIN assets a ON cgr.asset_id = a.id
            WHERE cgr.owner_id = ? AND cgr.financial_year = ?
            ORDER BY cgr.sale_date DESC
        """, (profile["person_id"], profile["financial_year"]))
        return [dict(row) for row in cursor.fetchall()]
    except HTTPException:
        raise
    finally:
        conn.close()

# income.frequency is free text in the generic entity form; these are the values the
# rest of the app uses (see Dashboard.tsx's FREQ_MULTIPLIER).
INCOME_FREQUENCY_PER_YEAR = {"monthly": 12, "quarterly": 4, "half-yearly": 2, "yearly": 1, "annually": 1, "annual": 1}

def _estimate_capital_gains(cursor, person_id: int, notes: Optional[list] = None) -> dict:
    """
    Unrealized gains on this person's Stock/Mutual Fund/Property assets, treating
    current_value as a hypothetical sale price today (distinct from the realized
    capital_gains_realizations log). Equity -> 111A/112A, property LTCG -> 112,
    property STCG -> slab rates.
    """
    gains = {"stcg_111a": 0.0, "ltcg_112a": 0.0, "ltcg_112": 0.0, "stcg_slab": 0.0}
    today_str = date.today().strftime("%Y-%m-%d")
    cursor.execute("""
        SELECT a.id, a.name, a.current_value, a.cost_price, a.purchase_date, a.metadata_json, c.name as category_name
        FROM assets a JOIN categories c ON a.category_id = c.id
        WHERE a.owner_id = ? AND c.name IN ('Stock', 'Mutual Fund', 'Property')
    """, (person_id,))
    notes = notes if notes is not None else []
    for a in cursor.fetchall():
        if a["category_name"] == "Property":
            if not a["purchase_date"]:
                continue
            gain = max(float(a["current_value"] or 0.0) - float(a["cost_price"] or 0.0), 0.0)
            key = "ltcg_112" if calc.determine_gain_type(a["purchase_date"], today_str, "Property") == "LTCG" else "stcg_slab"
            gains[key] += gain
            continue
        cursor.execute("SELECT * FROM holding_lots WHERE asset_id = ? ORDER BY purchase_date", (a["id"],))
        lots = [dict(l) for l in cursor.fetchall()]
        total_qty = sum(l["quantity"] for l in lots)
        if not lots or total_qty <= 0:
            continue
        # No recorded cost (e.g. the auto-migrated 1-unit, ₹0 opening lot): the whole value
        # would read as profit, so leave it out and say so.
        if sum(l["quantity"] * l["price_per_unit"] for l in lots) <= 0:
            notes.append(f"{a['name']} is left out of unrealised gains: its purchase lots have no cost recorded. Add them under Assets.")
            continue
        cg = calc.calculate_capital_gains_for_lots(
            lots, total_qty, float(a["current_value"] or 0.0) / total_qty, today_str, a["category_name"])
        kind = fund_type(risk_metadata(a["metadata_json"])) if a["category_name"] == "Mutual Fund" else None
        if kind and kind not in EQUITY_ORIENTED_FUND_TYPES:
            # Debt, liquid and other funds bought after 1 April 2023 are taxed at slab
            # rates, short or long term (older units may still qualify for 112; not modelled).
            gains["stcg_slab"] += max(cg["stcg_amount"], 0.0) + max(cg["ltcg_amount"], 0.0)
            continue
        gains["stcg_111a"] += max(cg["stcg_amount"], 0.0)
        gains["ltcg_112a"] += max(cg["ltcg_amount"], 0.0)
    return gains

def _tax_profile_to_engine_inputs(cursor, profile: dict, rules: dict) -> tuple:
    """Maps a stored tax profile (+ household income rows + itemised deductions) onto engine inputs."""
    notes = []

    salary_total = 0.0
    other = {"rental_gross": 0.0, "dividends": 0.0, "interest": 0.0, "other": 0.0}
    cursor.execute("""
        SELECT i.amount, i.frequency, c.name as category_name
        FROM income i LEFT JOIN categories c ON i.category_id = c.id
        WHERE i.taxable = 1 AND (i.owner_id = ? OR i.owner_id IS NULL)
    """, (profile["person_id"],))
    for row in cursor.fetchall():
        freq = (row["frequency"] or "Monthly").strip().lower()
        per_year = INCOME_FREQUENCY_PER_YEAR.get(freq)
        if per_year is None:
            per_year = 1
            notes.append(f"Income frequency '{row['frequency']}' not recognised; treated as an annual amount.")
        annual = float(row["amount"] or 0.0) * per_year
        category = row["category_name"]
        if category in ("Salary", "Bonus"):
            salary_total += annual
        elif category == "Rental":
            other["rental_gross"] += annual
        elif category == "Dividend":
            other["dividends"] += annual
        elif category == "Interest":
            other["interest"] += annual
        else:
            other["other"] += annual
    other["other"] += float(profile.get("other_income_reported") or 0.0)
    notes.append("Income is this person's taxable income entries (Income & Expenses).")

    basic = min(float(profile.get("hra_basic_salary") or 0.0), salary_total)
    hra = min(float(profile.get("hra_received") or 0.0), salary_total - basic)
    if float(profile.get("hra_basic_salary") or 0.0) + float(profile.get("hra_received") or 0.0) > salary_total + 0.5:
        notes.append("Basic + HRA entered on the profile exceed salary income from income rows; capped to salary income.")

    cursor.execute("SELECT section, COALESCE(SUM(amount), 0.0) as total FROM tax_deductions WHERE tax_profile_id = ? GROUP BY section", (profile["id"],))
    sections = {row["section"]: float(row["total"]) for row in cursor.fetchall()}

    # 80D entries aren't split between self/family and parents on the profile: fill
    # self/family up to its cap first, the remainder counts as non-senior parents.
    old_80d = rules["regimes"]["old"]["deductions"]["value"]["80D"]
    bands = rules["age_bands"]["value"]
    representative_age = {"Below60": bands["senior_from_age"] - 1, "60to80": bands["senior_from_age"],
                          "Above80": bands["super_senior_from_age"]}[profile.get("age_band") or "Below60"]
    self_cap = old_80d["self_family_senior"] if representative_age >= bands["senior_from_age"] else old_80d["self_family"]
    total_80d = sections.get("80D", 0.0)
    if total_80d > self_cap:
        notes.append("80D entries are not split by insured person; the amount above the self/family cap is treated as non-senior parents' premium.")

    savings_interest = min(sections.get("80TTA", 0.0), other["interest"])
    annual_rent = float(profile.get("hra_rent_paid") or 0.0)

    inputs = {
        "profile": {"age": representative_age, "metro": bool(profile.get("is_metro_city")),
                    "selected_regime": (profile.get("regime") or "").lower() or None},
        "salary": {"basic": basic, "hra": hra, "other_taxable_allowances": salary_total - basic - hra},
        "hra": {"rent_paid_monthly": annual_rent / 12.0, "months_rented": 12},
        "other_income": {"rental_gross": other["rental_gross"], "dividends": other["dividends"],
                         "interest_savings": savings_interest, "interest_deposits": other["interest"] - savings_interest,
                         "other": other["other"]},
        "capital_gains": _estimate_capital_gains(cursor, profile["person_id"], notes),
        "deductions": {
            "80C": {"other": sections.get("80C", 0.0)},
            "employee_nps": sections.get("80CCD1B", 0.0),
            "80D_self_family": min(total_80d, self_cap),
            "80D_parents": max(total_80d - self_cap, 0.0),
            "home_loan_interest_self_occupied": float(profile.get("home_loan_interest_self_occupied") or 0.0) + sections.get("24B", 0.0),
            "80E": sections.get("80E", 0.0),
            "80G_eligible": sections.get("80G", 0.0),
            "other_old": sections.get("Other", 0.0),
        },
        "prepaid": {"tds": float(profile.get("tds_deducted") or 0.0), "advance_tax": float(profile.get("advance_tax_paid") or 0.0)},
    }
    return inputs, notes

@app.get("/api/tax-profiles/{id}/estimate")
def get_tax_estimate(id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM tax_profiles WHERE id = ?", (id,))
        profile = cursor.fetchone()
        if not profile:
            raise HTTPException(status_code=404, detail="Tax profile not found")
        profile = dict(profile)
        rules, rule_meta = tax_store.resolve(conn, profile["financial_year"])
        inputs, notes = _tax_profile_to_engine_inputs(cursor, profile, rules)
        computed_at = datetime.now().isoformat(timespec="seconds")

        # Unrealized capital gains are a "what if I sold today" figure, so the regime
        # comparison excludes them; their incremental tax (including any surcharge
        # interplay) is reported separately and added to the estimated total.
        with_gains = tax_engine.compute(inputs, rules, computed_at)
        without_gains = tax_engine.compute({**inputs, "capital_gains": {}}, rules, computed_at)
        for computed in (with_gains, without_gains):
            computed["rule_set"].update({"name": rule_meta["name"], "source": rule_meta["source"],
                                         "content_hash": rule_meta["content_hash"]})
        old, new = without_gains["regimes"]["old"], without_gains["regimes"]["new"]
        recommended = without_gains["comparison"]["best_regime"]
        cg_tax = with_gains["regimes"][recommended]["total_tax"] - without_gains["regimes"][recommended]["total_tax"]
        cg = inputs["capital_gains"]
        recommended_total = without_gains["regimes"][recommended]["total_tax"] + cg_tax
        prepaid = old["prepaid_taxes"]

        def regime_summary(r: dict) -> dict:
            return {"gross_tax": r["tax_before_rebate"], "rebate_87a": r["rebate"], "surcharge": r["surcharge"],
                    "cess": r["cess"], "total_tax": r["total_tax"]}

        warnings = without_gains["warnings"] + [
            {"code": "profile_mapping", "severity": "info", "regime": None, "message": n} for n in notes]
        return {
            "financial_year": profile["financial_year"],
            "rule_set": without_gains["rule_set"],
            "gross_income": old["gross_total_income"],
            "standard_deduction": old["standard_deduction"],
            "standard_deduction_new": new["standard_deduction"],
            "total_80c_deductions": old["deductions"].get("80C", 0.0),
            "total_80d_deductions": old["deductions"].get("80D", 0.0),
            "other_deductions": round(old["deductions_total"] - old["deductions"].get("80C", 0.0) - old["deductions"].get("80D", 0.0), 2),
            "hra_exemption": old["hra"]["exemption"] if old["hra"] else 0.0,
            "taxable_income_old": old["total_income"],
            "taxable_income_new": new["total_income"],
            "capital_gains": {"stcg": round(cg["stcg_111a"] + cg["stcg_slab"], 2),
                              "ltcg": round(cg["ltcg_112a"] + cg["ltcg_112"], 2), "tax": round(cg_tax, 2)},
            "old_regime": regime_summary(old),
            "new_regime": regime_summary(new),
            "recommended_regime": recommended.capitalize(),
            "savings_with_recommended": abs(without_gains["comparison"]["savings_old_minus_new"]),
            "estimated_total_tax": round(recommended_total, 2),
            "prepaid_taxes": prepaid,
            "estimated_refund_or_due": round(prepaid - recommended_total, 2),
            "warnings": warnings,
            "computation": without_gains,
            "disclaimer": f"Estimate only, computed under rule set {rules['rule_set_id']} v{rules['version']}. Not tax advice — verify against the current Act and Finance Act before filing.",
        }
    except tax_store.RuleStoreError as e:
        raise HTTPException(status_code=e.status_code, detail=str(e))
    except tax_engine.TaxInputError as e:
        raise HTTPException(status_code=422, detail=f"Could not map tax profile to engine inputs: {e}")
    except HTTPException:
        raise
    finally:
        conn.close()

@app.delete("/api/metadata/{meta_type}/{id}")
def delete_metadata(meta_type: str, id: int):
    if meta_type not in ["categories", "institutions", "financial_rules", "people"]:
        raise HTTPException(status_code=400, detail="Invalid metadata type")
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(f"DELETE FROM {meta_type} WHERE id = ?", (id,))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- Metadata Configuration REST API ---
@app.get("/api/metadata/{meta_type}")
def get_metadata(meta_type: str):
    if meta_type not in ["categories", "institutions", "financial_rules", "people"]:
        raise HTTPException(status_code=400, detail="Invalid metadata type")
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        if meta_type == "institutions":
            cursor.execute("SELECT * FROM institutions ORDER BY name ASC")
        elif meta_type == "people":
            cursor.execute("SELECT * FROM people ORDER BY name ASC")
        else:
            cursor.execute(f"SELECT * FROM {meta_type}")
        rows = cursor.fetchall()
        result = []
        for r in rows:
            d = dict(r)
            if d.get("fields_definition"):
                try:
                    d["fields_definition"] = json.loads(d["fields_definition"])
                except:
                    pass
            result.append(d)
        return result
    finally:
        conn.close()

@app.post("/api/metadata/{meta_type}")
def create_metadata(meta_type: str, payload: dict):
    if meta_type not in ["categories", "institutions", "financial_rules", "people"]:
        raise HTTPException(status_code=400, detail="Invalid metadata type")
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        if payload.get("fields_definition") and not isinstance(payload["fields_definition"], str):
            payload["fields_definition"] = json.dumps(payload["fields_definition"])
        
        columns = list(payload.keys())
        placeholders = [f":{c}" for c in columns]
        cursor.execute(f"INSERT OR IGNORE INTO {meta_type} ({', '.join(columns)}) VALUES ({', '.join(placeholders)})", payload)
        conn.commit()
        return {"status": "success", "id": cursor.lastrowid}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.put("/api/metadata/{meta_type}/{id}")
def update_metadata(meta_type: str, id: int, payload: dict):
    if meta_type not in ["categories", "institutions", "people"]:
        raise HTTPException(status_code=400, detail="Invalid metadata type for update")
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        if payload.get("fields_definition") and not isinstance(payload["fields_definition"], str):
            payload["fields_definition"] = json.dumps(payload["fields_definition"])
            
        set_clauses = [f"{c} = :{c}" for c in payload.keys()]
        payload["id"] = id
        cursor.execute(f"UPDATE {meta_type} SET {', '.join(set_clauses)} WHERE id = :id", payload)
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- Rules Router ---
@app.put("/api/rules")
def update_rule(rule: RuleSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "UPDATE financial_rules SET value=?, is_enabled=? WHERE rule_key=?",
            (rule.value, int(rule.is_enabled), rule.rule_key)
        )
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- AI Assistant ---
# Explain-only, never calculate: every prompt embeds numbers already computed by
# calculations.py/decision_engine.py; the model's reply is always returned as
# opaque display text, never parsed as instructions or written back to any
# table. Off by default (ai_settings.provider = 'none') — no outbound network
# call happens anywhere in the app until a user explicitly configures a
# provider in Settings. See Architecture/07-AI-Architecture.md.

class AISettingsSchema(BaseModel):
    provider: str = "none"
    endpoint: Optional[str] = ""
    api_key: Optional[str] = ""
    model: Optional[str] = ""
    max_tokens: Optional[int] = 500

class AIExplainRequest(BaseModel):
    rule_key: Optional[str] = None
    recommendation_index: Optional[int] = None
    context_note: Optional[str] = ""

class AIAskRequest(BaseModel):
    question: str

@app.get("/api/ai/settings")
def get_ai_settings():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM ai_settings WHERE id = 1")
        row = cursor.fetchone()
        result = dict(row) if row else {"provider": "none", "max_tokens": 500}
        result["api_key_set"] = bool(result.get("api_key"))
        result.pop("api_key", None)
        return result
    finally:
        conn.close()

@app.put("/api/ai/settings")
def update_ai_settings(settings: AISettingsSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT api_key FROM ai_settings WHERE id = 1")
        existing = cursor.fetchone()
        # Blank api_key in the incoming payload means "leave unchanged", so
        # re-saving provider/model doesn't silently wipe a previously-stored key.
        api_key = settings.api_key if settings.api_key else (existing["api_key"] if existing else None)
        cursor.execute("""
            INSERT INTO ai_settings (id, provider, endpoint, api_key, model, max_tokens, updated_at)
            VALUES (1, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(id) DO UPDATE SET
                provider=excluded.provider, endpoint=excluded.endpoint, api_key=excluded.api_key,
                model=excluded.model, max_tokens=excluded.max_tokens, updated_at=excluded.updated_at
        """, (settings.provider, settings.endpoint, api_key, settings.model, settings.max_tokens, datetime.now().isoformat()))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

@app.post("/api/ai/test")
def test_ai_connection():
    from backend.ai.providers import get_provider, AIProviderError
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM ai_settings WHERE id = 1")
        row = cursor.fetchone()
        if not row or row["provider"] == "none":
            raise HTTPException(status_code=400, detail="No AI provider configured")
        provider = get_provider(dict(row))
        reply = provider.complete("You are a connectivity test.", "Reply with the single word OK.", max_tokens=10)
        return {"status": "success", "reply": reply.strip()}
    except AIProviderError as e:
        raise HTTPException(status_code=502, detail=str(e))
    except HTTPException:
        raise
    finally:
        conn.close()

@app.post("/api/ai/explain")
def ai_explain(req: AIExplainRequest):
    from backend.ai.providers import get_provider, AIProviderError
    from backend.ai.context import build_explain_context
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM ai_settings WHERE id = 1")
        row = cursor.fetchone()
        if not row or row["provider"] == "none":
            raise HTTPException(status_code=400, detail="AI Assistant is not configured")
        provider = get_provider(dict(row))
        system_prompt, user_prompt = build_explain_context(conn, req.rule_key, req.recommendation_index, req.context_note)
        try:
            reply = provider.complete(system_prompt, user_prompt, max_tokens=row["max_tokens"] or 500)
        except AIProviderError as e:
            raise HTTPException(status_code=502, detail=str(e))
        return {"explanation": reply}
    except HTTPException:
        raise
    finally:
        conn.close()

@app.post("/api/ai/ask")
def ai_ask(req: AIAskRequest):
    from backend.ai.providers import get_provider, AIProviderError
    from backend.ai.context import build_ask_context
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM ai_settings WHERE id = 1")
        row = cursor.fetchone()
        if not row or row["provider"] == "none":
            raise HTTPException(status_code=400, detail="AI Assistant is not configured")
        provider = get_provider(dict(row))
        system_prompt, user_prompt = build_ask_context(conn, req.question)
        try:
            reply = provider.complete(system_prompt, user_prompt, max_tokens=row["max_tokens"] or 500)
        except AIProviderError as e:
            raise HTTPException(status_code=502, detail=str(e))
        return {"answer": reply}
    except HTTPException:
        raise
    finally:
        conn.close()

# --- Snapshots Router ---
@app.get("/api/snapshots")
def get_snapshots():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM snapshots ORDER BY month ASC")
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()

@app.post("/api/snapshots")
def create_snapshot(month: str = Form(...), snapshot_data_str: str = Form(...)):
    """
    Saves a monthly snapshot. The frontend submits a snapshot month YYYY-MM
    and a JSON representation of all current financial objects.
    """
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        data = json.loads(snapshot_data_str)
        
        # Calculate summary metrics for snapshot
        total_liquid = sum(float(a.get("current_balance") if a.get("current_balance") is not None else 0.0) for a in data.get("accounts", []))
        total_assets_val = sum(float(a.get("current_value") if a.get("current_value") is not None else 0.0) for a in data.get("assets", []))
        total_debt = sum(float(l.get("outstanding") if l.get("outstanding") is not None else 0.0) for l in data.get("liabilities", []))
        
        total_assets = total_liquid + total_assets_val
        net_worth = total_assets - total_debt
        
        cursor.execute(
            """INSERT INTO snapshots (month, net_worth, liquid_funds, total_assets, total_liabilities, snapshot_data)
               VALUES (?, ?, ?, ?, ?, ?)
               ON CONFLICT(month) DO UPDATE SET
               net_worth=excluded.net_worth,
               liquid_funds=excluded.liquid_funds,
               total_assets=excluded.total_assets,
               total_liabilities=excluded.total_liabilities,
               snapshot_data=excluded.snapshot_data""",
            (month, net_worth, total_liquid, total_assets, total_debt, snapshot_data_str)
        )
        conn.commit()
        return {"status": "success", "net_worth": net_worth}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    finally:
        conn.close()

# --- Simulations Engine ---
@app.post("/api/simulations/run")
def run_simulation(req: SimulationRequestSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        # Load current state
        cursor.execute("""
            SELECT COALESCE(SUM(
                (SELECT COALESCE(SUM(t.amount), 0.0) FROM transactions t WHERE t.dest_ref_type = 'Account' AND t.dest_ref_id = a.id) -
                (SELECT COALESCE(SUM(t.amount), 0.0) FROM transactions t WHERE t.source_ref_type = 'Account' AND t.source_ref_id = a.id)
            ), 0.0) AS total_liquid
            FROM accounts a
        """)
        total_liquid = cursor.fetchone()["total_liquid"]
        
        cursor.execute("SELECT id, name, interest_rate as roi, emi, outstanding FROM liabilities")
        loans = [dict(row) for row in cursor.fetchall()]
        
        cursor.execute("""
            SELECT current_value FROM assets 
            WHERE category_id != (SELECT id FROM categories WHERE name = 'Property')
        """)
        total_investments = sum(row["current_value"] for row in cursor.fetchall())
        
        cursor.execute("""
            SELECT current_value FROM assets 
            WHERE category_id = (SELECT id FROM categories WHERE name = 'Property')
        """)
        total_properties = sum(row["current_value"] for row in cursor.fetchall())
        
        # See /api/init's goals query for why the alias is not named current_amount.
        cursor.execute("""
            SELECT
                g.*,
                CASE WHEN EXISTS (SELECT 1 FROM goal_holdings gh WHERE gh.goal_id = g.id)
                    THEN COALESCE(
                        (
                            SELECT SUM(a.current_value * gh.allocation_weight)
                            FROM assets a
                            JOIN goal_holdings gh ON a.id = gh.asset_id
                            WHERE gh.goal_id = g.id
                        ), 0.0
                    )
                    ELSE g.current_amount
                END AS derived_current_amount
            FROM goals g
        """)
        goals = []
        for row in cursor.fetchall():
            d = dict(row)
            d["current_amount"] = d.pop("derived_current_amount")
            goals.append(d)

        total_assets = total_liquid + total_investments + total_properties
        total_debt = sum(l["outstanding"] for l in loans)
        current_net_worth = total_assets - total_debt
        current_monthly_emi = sum(l["emi"] for l in loans)
        
        # Calculate current debt free date (max tenure of active loans)
        # Debt-free date: loans with an EMI. Cards and other debts without one have no
        # schedule, and would otherwise read as never paid off.
        current_tenures = [calc.calculate_remaining_tenure(l["outstanding"], l["roi"], l["emi"]) for l in loans if l["outstanding"] > 0 and (l["emi"] or 0) > 0]
        max_tenure_current = max(current_tenures) if current_tenures else 0.0
        
        # Cloned simulated variables
        sim_liquid = total_liquid
        sim_investments = total_investments
        sim_properties = total_properties
        sim_loans = [dict(l) for l in loans]
        sim_goals = [dict(g) for g in goals]
        sim_salary_surplus = 0.0
        
        recommendation = ""
        
        # Apply Scenario
        s_type = req.scenario_type
        params = req.params
        
        if s_type == "close_loan":
            loan_id = params.get("loan_id")
            # Find the loan
            loan = next((l for l in sim_loans if l["id"] == loan_id), None)
            if loan and loan["outstanding"] > 0:
                outstanding_cost, old_emi = loan["outstanding"], loan["emi"]
                # Work out what closing saves before zeroing the loan (it used to zero the
                # EMI first, so both figures always read ₹0).
                interest_saved = calc.calculate_remaining_interest(outstanding_cost, loan["roi"], old_emi) if old_emi > 0 else 0.0
                sim_liquid -= outstanding_cost
                loan["outstanding"] = 0.0
                loan["emi"] = 0.0
                recommendation = (f"Closing {loan['name']} now ({inr(outstanding_cost)}) saves about {inr(interest_saved)} of future interest "
                                  f"and frees {inr(old_emi)} a month.")
                
        elif s_type == "delay_goal":
            goal_id = params.get("goal_id")
            delay_months = params.get("delay_months", 12)
            goal = next((g for g in sim_goals if g["id"] == goal_id), None)
            if goal:
                # Simulating delay reduces required monthly contribution to reach target
                gap = max(goal["target_amount"] - goal["current_amount"], 0.0)
                old_contribution = goal["monthly_contribution"]
                current_rem_months = gap / old_contribution if old_contribution > 0 else 12.0
                new_rem_months = current_rem_months + delay_months
                new_contribution = gap / new_rem_months if new_rem_months > 0 else 0.0
                goal["monthly_contribution"] = new_contribution
                recommendation = (f"Delaying '{goal['name']}' by {delay_months} months lowers what it needs each month to "
                                  f"{inr(new_contribution)}, {inr(max(old_contribution - new_contribution, 0.0))} less than now.")
                
        elif s_type == "increase_sip":
            increment = params.get("increment_amount", 5000.0)
            sim_investments += (increment * 12.0) # simulation over 1 year
            recommendation = f"Increasing your SIP by {inr(increment)}/month yields a portfolio size increase of {inr(increment*12)} in 1 year (excluding compounding returns)."
            
        elif s_type == "salary_change":
            change = params.get("change_amount", 10000.0)
            sim_liquid += (change * 12.0)
            surplus_msg = "surplus" if change >= 0 else "deficit"
            recommendation = f"A salary change of {inr(change)}/month creates a yearly {surplus_msg} of {inr(change*12)} that can be used for prepayments."
            
        elif s_type == "receive_bonus":
            bonus = params.get("bonus_amount", 100000.0)
            allocation_target = params.get("allocation_target", "liquid") # liquid, prepay, invest
            
            if allocation_target == "liquid":
                sim_liquid += bonus
                recommendation = f"Receiving {inr(bonus)} as bonus increases your cash liquidity instantly."
            elif allocation_target == "prepay":
                # Prepay highest interest loan
                active_loans = [l for l in sim_loans if l["outstanding"] > 0]
                if active_loans:
                    active_loans.sort(key=lambda x: x["roi"], reverse=True)
                    target_loan = active_loans[0]
                    prepay_amt = min(bonus, target_loan["outstanding"])
                    interest_saved = calc.calculate_interest_saved_after_prepayment(target_loan["outstanding"], target_loan["roi"], target_loan["emi"], prepay_amt)
                    target_loan["outstanding"] -= prepay_amt
                    recommendation = f"Allocating your bonus of {inr(bonus)} to prepay {target_loan['name']} saves {inr(interest_saved)} in remaining interest."
                else:
                    sim_liquid += bonus
                    recommendation = "No active loans to prepay. Bonus added to liquidity."
            elif allocation_target == "invest":
                sim_investments += bonus
                recommendation = f"Investing your bonus of {inr(bonus)} increases compound growth capital."
                
        elif s_type == "roi_reduction":
            loan_id = params.get("loan_id")
            new_roi = params.get("new_roi", 7.0)
            loan = next((l for l in sim_loans if l["id"] == loan_id), None)
            if loan and loan["outstanding"] > 0:
                old_interest = calc.calculate_remaining_interest(loan["outstanding"], loan["roi"], loan["emi"])
                new_interest = calc.calculate_remaining_interest(loan["outstanding"], new_roi, loan["emi"])
                saved = old_interest - new_interest
                loan["roi"] = new_roi
                recommendation = f"Refinancing {loan['name']} to {new_roi}% ROI saves {inr(saved)} in interest."

        elif s_type in ["brother_property", "large_purchase"]:
            purchase_name = params.get("purchase_name", "Family Property")
            cost = float(params.get("property_cost") or params.get("cost") or 1500000.0)
            funding_source = params.get("funding_source", "liquid") # liquid, loan
            
            if funding_source == "liquid":
                sim_liquid -= cost
                recommendation = f"Funding your '{purchase_name}' of {inr(cost)} directly reduces liquidity, lowering net worth by the same amount."
            else:
                roi = float(params.get("roi") or 8.5)
                tenure_months = int(params.get("tenure_months") or 120)
                calculated_emi = calc.calculate_monthly_emi(cost, roi, tenure_months)
                
                # Add a simulated loan
                sim_loans.append({
                    "id": -99,
                    "name": f"Simulated {purchase_name} Loan",
                    "lender": "Bank",
                    "outstanding": cost,
                    "roi": roi,
                    "emi": calculated_emi,
                    "tenure": tenure_months
                })
                recommendation = f"Taking a loan of {inr(cost)} at {roi}% ROI to fund '{purchase_name}' increases debt by {inr(cost)} and adds a monthly EMI commitment of {inr(calculated_emi)} for {tenure_months} months."

        # Recalculate simulated outcomes
        sim_assets = sim_liquid + sim_investments + sim_properties
        sim_debt = sum(l["outstanding"] for l in sim_loans)
        sim_net_worth = sim_assets - sim_debt
        sim_monthly_emi = sum(l["emi"] for l in sim_loans)
        
        sim_tenures = [calc.calculate_remaining_tenure(l["outstanding"], l["roi"], l["emi"]) for l in sim_loans if l["outstanding"] > 0 and (l["emi"] or 0) > 0]
        if sim_liquid < 0 <= total_liquid:
            recommendation += f" You'd need {inr(-sim_liquid)} more cash than your accounts hold."
        max_tenure_sim = max(sim_tenures) if sim_tenures else 0.0
        
        # Handle JSON compliance for infinite tenures
        current_tenure_val = -1.0 if math.isinf(max_tenure_current) else round(max_tenure_current, 1)
        simulated_tenure_val = -1.0 if math.isinf(max_tenure_sim) else round(max_tenure_sim, 1)
        
        if math.isinf(max_tenure_sim) or math.isinf(max_tenure_current):
            diff_tenure_val = -1.0
        else:
            diff_tenure_val = round(max_tenure_sim - max_tenure_current, 1)
        
        return {
            "current": {
                "net_worth": current_net_worth,
                "total_debt": total_debt,
                "liquid_funds": total_liquid,
                "monthly_emi": current_monthly_emi,
                "remaining_tenure_months": current_tenure_val
            },
            "simulated": {
                "net_worth": sim_net_worth,
                "total_debt": sim_debt,
                "liquid_funds": sim_liquid,
                "monthly_emi": sim_monthly_emi,
                "remaining_tenure_months": simulated_tenure_val
            },
            "difference": {
                "net_worth": sim_net_worth - current_net_worth,
                "total_debt": sim_debt - total_debt,
                "liquid_funds": sim_liquid - total_liquid,
                "monthly_emi": sim_monthly_emi - current_monthly_emi,
                "tenure_months": diff_tenure_val
            },
            "recommendation": recommendation
        }
    finally:
        conn.close()

@app.post("/api/simulations/debt_payoff")
def run_debt_payoff_simulation(req: DebtPayoffRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT name, outstanding, interest_rate as roi, emi FROM liabilities WHERE outstanding > 0")
        rows = cursor.fetchall()
        
        loans_list = [
            {
                "name": row["name"],
                "outstanding": float(row["outstanding"]),
                "roi": float(row["roi"]),
                "emi": float(row["emi"])
            }
            for row in rows
        ]
        
        if not loans_list:
            return {
                "baseline": {"total_interest": 0.0, "months": 0, "schedule": []},
                "optimized": {"total_interest": 0.0, "months": 0, "schedule": []}
            }
            
        baseline = calc.simulate_multi_loan_payoff(
            loans_list=json.loads(json.dumps(loans_list)),
            extra_monthly=0.0,
            strategy="baseline",
            prepayment_frequency="none",
            prepayment_amount=0.0
        )
        
        optimized = calc.simulate_multi_loan_payoff(
            loans_list=json.loads(json.dumps(loans_list)),
            extra_monthly=req.extra_monthly,
            strategy=req.strategy,
            prepayment_frequency=req.prepayment_frequency,
            prepayment_amount=req.prepayment_amount,
            prepayment_month=req.prepayment_month
        )
        
        return {
            "baseline": baseline,
            "optimized": optimized
        }
    finally:
        conn.close()

@app.post("/api/simulations/net_worth_forecast")
def run_net_worth_forecast(req: NetWorthForecastRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            SELECT COALESCE(SUM(
                (SELECT COALESCE(SUM(t.amount), 0.0) FROM transactions t WHERE t.dest_ref_type = 'Account' AND t.dest_ref_id = a.id) -
                (SELECT COALESCE(SUM(t.amount), 0.0) FROM transactions t WHERE t.source_ref_type = 'Account' AND t.source_ref_id = a.id)
            ), 0.0) AS total_liquid
            FROM accounts a
        """)
        liquid = cursor.fetchone()["total_liquid"]
        cursor.execute("SELECT current_value FROM assets")
        assets_val = sum(row["current_value"] for row in cursor.fetchall())
        initial_assets = liquid + assets_val
        
        cursor.execute("SELECT name, outstanding, interest_rate, emi FROM liabilities WHERE outstanding > 0")
        liabs_raw = cursor.fetchall()
        liabilities = [dict(row) for row in liabs_raw]
        initial_debt = sum(l["outstanding"] for l in liabilities)
        
        sim_loans = [
            {
                "outstanding": float(l["outstanding"]),
                "roi": float(l["interest_rate"]),
                "emi": float(l["emi"])
            }
            for l in liabilities
        ]
        
        schedule = []
        current_assets = initial_assets
        annual_savings = req.monthly_savings * 12.0
        
        schedule.append({
            "year": 0,
            "assets": round(current_assets, 2),
            "debt": round(initial_debt, 2),
            "net_worth_nominal": round(current_assets - initial_debt, 2),
            "net_worth_real": round(current_assets - initial_debt, 2)
        })
        
        for y in range(1, req.horizon_years + 1):
            current_assets = current_assets * (1.0 + req.expected_return / 100.0) + annual_savings * ((1.0 + req.inflation_rate / 100.0) ** (y - 1))
            
            for _ in range(12):
                for loan in sim_loans:
                    if loan["outstanding"] > 0:
                        interest = loan["outstanding"] * (loan["roi"] / 100.0) / 12.0
                        payment = min(loan["emi"], loan["outstanding"] + interest)
                        principal = payment - interest
                        loan["outstanding"] = max(loan["outstanding"] - principal, 0.0)
                        
            current_debt = sum(l["outstanding"] for l in sim_loans)
            nominal_nw = current_assets - current_debt
            inflation_factor = (1.0 + req.inflation_rate / 100.0) ** y
            real_nw = nominal_nw / inflation_factor
            
            schedule.append({
                "year": y,
                "assets": round(current_assets, 2),
                "debt": round(current_debt, 2),
                "net_worth_nominal": round(nominal_nw, 2),
                "net_worth_real": round(real_nw, 2)
            })
            
        return {
            "schedule": schedule,
            "final_net_worth_nominal": round(schedule[-1]["net_worth_nominal"], 2),
            "final_net_worth_real": round(schedule[-1]["net_worth_real"], 2),
            "initial_net_worth": round(initial_assets - initial_debt, 2)
        }
    finally:
        conn.close()

# --- MMex Data Import ---
@app.post("/api/import/mmex")
async def import_mmex_csv(file: UploadFile = File(...)):
    """
    Parses a Money Manager Ex (MMex) Accounts CSV file to import / sync account balances.
    Assumed CSV Columns: AccountName, AccountType, CurrentBalance, Notes, etc.
    """
    contents = await file.read()
    import pandas as pd
    import io
    
    try:
        df = pd.read_csv(io.BytesIO(contents))
        
        # Verify columns or attempt flexible lookup
        name_col = next((c for c in df.columns if 'name' in c.lower() or 'account' in c.lower()), None)
        bal_col = next((c for c in df.columns if 'balance' in c.lower() or 'current' in c.lower() or 'bal' in c.lower()), None)
        inst_col = next((c for c in df.columns if 'inst' in c.lower() or 'bank' in c.lower() or 'type' in c.lower()), "MMex Imported")
        
        if not name_col or not bal_col:
            raise HTTPException(status_code=400, detail="Could not identify Account Name or Balance columns in CSV. Columns: " + ", ".join(df.columns))
        
        conn = get_db_connection()
        cursor = conn.cursor()
        
        imported_count = 0
        for _, row in df.iterrows():
            name = str(row[name_col]).strip()
            val = float(row[bal_col])
            institution = str(row[inst_col]).strip() if isinstance(inst_col, str) and inst_col in df.columns else "MMex"
            
            # Resolve institution_id
            cursor.execute("SELECT id FROM institutions WHERE name = ?", (institution,))
            inst_row = cursor.fetchone()
            if inst_row:
                inst_id = inst_row["id"]
            else:
                cursor.execute("INSERT INTO institutions (name) VALUES (?)", (institution,))
                inst_id = cursor.lastrowid
                
            cursor.execute(
                """INSERT INTO accounts (name, institution_id, owner_id, status)
                   VALUES (?, ?, 1, 'Active')
                   ON CONFLICT(name, institution_id, owner_id) DO UPDATE SET status = 'Active'""",
                (name, inst_id)
            )
            account_id = cursor.lastrowid
            if not account_id:
                cursor.execute("SELECT id FROM accounts WHERE name = ? AND institution_id = ? AND owner_id = 1", (name, inst_id))
                account_id = cursor.fetchone()["id"]
                
            if val > 0.0:
                cursor.execute("""
                INSERT INTO transactions (source_ref_type, source_ref_id, dest_ref_type, dest_ref_id, type, amount, transaction_date, notes)
                VALUES ('None', NULL, 'Account', ?, 'Credit', ?, ?, 'CSV Import Initial Balance')
                """, (account_id, val, date.today().isoformat()))
            imported_count += 1
            
        conn.commit()
        conn.close()
        return {"status": "success", "imported": imported_count}
    except Exception as e:
        raise HTTPException(status_code=400, detail="Parsing error: " + str(e))

# --- Database Backup & Restore (logic in backend/backups.py) ---
def _backup_call(fn, *args):
    try:
        return fn(*args)
    except backups.BackupError as e:
        raise HTTPException(status_code=e.status_code, detail=str(e))

@app.post("/api/backup/create")
def create_backup():
    made = _backup_call(backups.create_backup, "manual")
    return {"status": "success", **made}

@app.post("/api/backup/restore")
def restore_backup(backup_name: str = Form(...)):
    return _backup_call(backups.restore_backup, backup_name)

@app.get("/api/backup/list")
def list_backups():
    # Names only, newest first (Settings.tsx); details at /api/backup/details.
    return [b["name"] for b in backups.list_backups()]

@app.get("/api/backup/details")
def list_backup_details():
    return {"folder": backups.backup_dir(), "backups": backups.list_backups()}

@app.api_route("/api/backup/auto-run", methods=["GET", "POST"])
def auto_run_backup(trigger: str = "open"):
    # Called by Electron while the window closes: report failures in the body (200), never
    # as an error status, so quitting is never held up by a failed backup.
    try:
        return backups.auto_backup(trigger)
    except Exception as e:
        return {"status": "error", "message": str(e)}

# --- Database Management Settings ---
@app.get("/api/settings/database")
def get_database_settings():
    from backend.db import load_settings, get_default_db_dir
    settings = load_settings()
    return {
        "active_path": get_db_path(),
        "default_path": os.path.join(get_default_db_dir(), "finance.db"),
        "recent_databases": settings.get("recent_databases", [get_db_path()]),
        "backup_on_open": settings.get("backup_on_open", True),
        "backup_on_close": settings.get("backup_on_close", True),
        "backup_count_limit": settings.get("backup_count_limit", 10),
        "custom_backup_dir": settings.get("custom_backup_dir", "")
    }

class BackupSettingsRequest(BaseModel):
    backup_on_open: bool
    backup_on_close: bool
    backup_count_limit: int
    custom_backup_dir: str

@app.post("/api/settings/backup")
def post_backup_settings(req: BackupSettingsRequest):
    from backend.db import load_settings, save_settings
    try:
        settings = load_settings()
        settings["backup_on_open"] = req.backup_on_open
        settings["backup_on_close"] = req.backup_on_close
        settings["backup_count_limit"] = req.backup_count_limit
        settings["custom_backup_dir"] = req.custom_backup_dir
        save_settings(settings)
        return {"status": "success"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

class DatabaseSwitchRequest(BaseModel):
    path: str

@app.post("/api/settings/database")
def post_switch_database(req: DatabaseSwitchRequest):
    from backend.db import switch_database
    try:
        switch_database(req.path)
        return {"status": "success", "active_path": req.path}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# --- Advanced EMI & Prepayment Calculator ---
class AdvancedEmiRequest(BaseModel):
    loan_amount: float
    roi: float
    tenure_months: int
    start_date: str
    prepayments: list = []

@app.post("/api/calculations/advanced-emi")
def get_advanced_emi_calculation(req: AdvancedEmiRequest):
    try:
        result = calc.calculate_advanced_emi_schedule(
            req.loan_amount,
            req.roi,
            req.tenure_months,
            req.start_date,
            req.prepayments
        )
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/calculations/advanced-emi/export/excel")
def export_emi_excel(req: AdvancedEmiRequest):
    try:
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        result = calc.calculate_advanced_emi_schedule(
            req.loan_amount,
            req.roi,
            req.tenure_months,
            req.start_date,
            req.prepayments
        )
        
        wb = Workbook()
        ws = wb.active
        ws.title = "EMI Amortization Schedule"
        
        header_font = Font(name='Segoe UI', size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid")
        title_font = Font(name='Segoe UI', size=16, bold=True)
        bold_font = Font(name='Segoe UI', size=11, bold=True)
        regular_font = Font(name='Segoe UI', size=11)
        align_right = Alignment(horizontal='right')
        align_left = Alignment(horizontal='left')
        align_center = Alignment(horizontal='center')
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        
        # Merge cells for title block
        ws.merge_cells("A1:H1")
        ws["A1"] = "FinanceOS - Advanced EMI Amortization Schedule"
        ws["A1"].font = title_font
        ws.row_dimensions[1].height = 25
        
        # Add basic parameters metadata
        ws["A3"] = "Loan Amount:"
        ws["B3"] = req.loan_amount
        ws["A4"] = "Interest Rate (ROI):"
        ws["B4"] = f"{req.roi}%"
        ws["A5"] = "Original Tenure (Months):"
        ws["B5"] = req.tenure_months
        ws["A6"] = "Start Date:"
        ws["B6"] = req.start_date
        
        for r in range(3, 7):
            ws[f"A{r}"].font = bold_font
            ws[f"B{r}"].font = regular_font
            
        ws["D3"] = "Monthly EMI:"
        ws["E3"] = result["monthly_emi"]
        ws["D4"] = "Total Prepayments:"
        ws["E4"] = result["total_prepayments"]
        ws["D5"] = "Total Interest:"
        ws["E5"] = result["total_interest"]
        ws["D6"] = "Actual Closure Tenure:"
        ws["E6"] = f"{result['tenure_months']} Months"
        
        for r in range(3, 7):
            ws[f"D{r}"].font = bold_font
            ws[f"E{r}"].font = regular_font
            
        # Add table headers
        headers = ["Month", "Payment Date", "Beginning Balance", "EMI", "Interest Paid", "Principal Paid", "Prepayment", "Ending Balance"]
        ws.append([]) # Row 7 empty
        ws.append(headers) # Row 8
        
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=8, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center if col_idx <= 2 else align_right
            
        ws.row_dimensions[8].height = 24
        
        # Append monthly rows
        for idx, row in enumerate(result["schedule"], 9):
            ws.append([
                row["month"],
                row["date"],
                row["beginning_balance"],
                row["emi"],
                row["interest"],
                row["principal"],
                row["prepayments"],
                row["ending_balance"]
            ])
            ws.cell(row=idx, column=1).alignment = align_center
            ws.cell(row=idx, column=2).alignment = align_center
            for c in range(1, 9):
                cell = ws.cell(row=idx, column=c)
                cell.font = regular_font
                cell.border = thin_border
                if c >= 3:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = align_right
                    
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        headers = {
            'Content-Disposition': 'attachment; filename="emi_amortization_schedule.xlsx"'
        }
        return StreamingResponse(
            output, 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# --- EMI Saved Profiles CRUD Routers ---
@app.get("/api/settings/emi-profiles")
def get_emi_profiles():
    from backend.db import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, name, loan_amount, roi, tenure_value, tenure_unit, start_date, prepayments FROM emi_profiles")
    rows = cursor.fetchall()
    conn.close()
    
    profiles = []
    for r in rows:
        profiles.append({
            "id": r["id"],
            "name": r["name"],
            "loan_amount": r["loan_amount"],
            "roi": r["roi"],
            "tenure_value": r["tenure_value"],
            "tenure_unit": r["tenure_unit"],
            "start_date": r["start_date"],
            "prepayments": json.loads(r["prepayments"])
        })
    return profiles

class EmiProfileRequest(BaseModel):
    name: str
    loan_amount: float
    roi: float
    tenure_value: int
    tenure_unit: str
    start_date: str
    prepayments: list = []

@app.post("/api/settings/emi-profiles")
def save_emi_profile(req: EmiProfileRequest):
    from backend.db import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT id FROM emi_profiles WHERE name = ?", (req.name,))
        row = cursor.fetchone()
        prepay_json = json.dumps(req.prepayments)
        if row:
            cursor.execute("""
                UPDATE emi_profiles 
                SET loan_amount = ?, roi = ?, tenure_value = ?, tenure_unit = ?, start_date = ?, prepayments = ?
                WHERE name = ?
            """, (req.loan_amount, req.roi, req.tenure_value, req.tenure_unit, req.start_date, prepay_json, req.name))
        else:
            cursor.execute("""
                INSERT INTO emi_profiles (name, loan_amount, roi, tenure_value, tenure_unit, start_date, prepayments)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, (req.loan_amount, req.roi, req.tenure_value, req.tenure_unit, req.start_date, prepay_json, req.name) if row else (req.name, req.loan_amount, req.roi, req.tenure_value, req.tenure_unit, req.start_date, prepay_json))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/settings/emi-profiles/{profile_id}")
def delete_emi_profile(profile_id: int):
    from backend.db import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM emi_profiles WHERE id = ?", (profile_id,))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
 # --- Financial Calendar Endpoint ---
@app.get("/api/calendar")
def get_financial_calendar():
    from datetime import datetime, date
    import calendar
    from backend.db import get_db_connection
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    # Fetch loans (liabilities where interest_rate > 0 and emi > 0 and outstanding > 0)
    cursor.execute("""
        SELECT name, emi, outstanding, metadata_json 
        FROM liabilities 
        WHERE outstanding > 0 AND emi > 0 AND interest_rate > 0
    """)
    loans_raw = cursor.fetchall()
    loans = []
    for r in loans_raw:
        d = dict(r)
        meta = json.loads(d["metadata_json"]) if d.get("metadata_json") else {}
        # The EMI's debit day: the loan form's "EMI debit day", else the day of a legacy
        # start_date. With neither, the day is unknown (it used to fall back to today's
        # date, so every EMI moved day to day).
        day = meta.get("emi_day")
        if not day and meta.get("start_date"):
            try:
                day = datetime.strptime(meta["start_date"], "%Y-%m-%d").day
            except ValueError:
                day = None
        d["emi_day"] = int(day) if day and 1 <= int(day) <= 31 else None
        loans.append(d)
        
    # Fetch insurance policies (assets where category is 'Insurance')
    cursor.execute("""
        SELECT inst.name as company, ins.policy_number, ins.premium, ins.frequency, ins.renewal_date
        FROM assets a
        LEFT JOIN institutions inst ON a.institution_id = inst.id
        LEFT JOIN insurance_details ins ON ins.asset_id = a.id
        WHERE a.category_id = (SELECT id FROM categories WHERE name = 'Insurance')
    """)
    insurance_raw = cursor.fetchall()
    insurance = []
    for r in insurance_raw:
        d = dict(r)
        insurance.append({
            "company": d["company"] or "Insurance",
            "policy_number": d["policy_number"] or "N/A",
            "premium": float(d["premium"] or 0.0),
            "frequency": d["frequency"] or "Yearly",
            "renewal_date": d["renewal_date"]
        })
        
    # Fetch goals still needing contributions. current_amount is derived from linked
    # holdings when any exist (matching /api/init), else the stored manual value —
    # otherwise a goal already funded via linked assets keeps generating monthly
    # SIP reminder events forever off its stale stored amount.
    cursor.execute("""
        SELECT name, monthly_contribution, target_amount,
            CASE WHEN EXISTS (SELECT 1 FROM goal_holdings gh WHERE gh.goal_id = g.id)
                THEN COALESCE(
                    (
                        SELECT SUM(a.current_value * gh.allocation_weight)
                        FROM assets a
                        JOIN goal_holdings gh ON a.id = gh.asset_id
                        WHERE gh.goal_id = g.id
                    ), 0.0
                )
                ELSE g.current_amount
            END AS current_amount
        FROM goals g
        WHERE monthly_contribution > 0
    """)
    goals = [gr for gr in cursor.fetchall() if (gr["target_amount"] - gr["current_amount"]) > 0]

    # Fetch credit cards (liabilities where category is 'Credit Card')
    cursor.execute("""
        SELECT l.name, l.outstanding, ccd.statement_balance, ccd.due_date, ccd.statement_date, inst.name as issuer
        FROM liabilities l
        LEFT JOIN institutions inst ON l.lender_institution_id = inst.id
        LEFT JOIN credit_card_details ccd ON ccd.liability_id = l.id
        WHERE l.category_id = (SELECT id FROM categories WHERE name = 'Credit Card') AND l.outstanding > 0
    """)
    cc_raw = cursor.fetchall()
    credit_cards = []
    for r in cc_raw:
        d = dict(r)
        credit_cards.append({
            "name": d["name"],
            "outstanding": d["outstanding"],
            "statement_balance": float(d["statement_balance"] or 0.0),
            "due_date": int(d["due_date"] or 15),
            "statement_date": int(d["statement_date"] or 1),
            "issuer": d["issuer"] or "Issuer"
        })
        
    conn.close()
    
    events = []
    today = date.today()
    
    # Generate calendar events for the next 12 months
    for m_offset in range(12):
        target_month_val = today.month + m_offset
        target_year = today.year + (target_month_val - 1) // 12
        target_month = (target_month_val - 1) % 12 + 1
        
        # A. Loan EMIs
        for l in loans:
            if l["emi_day"]:
                due_date = date(target_year, target_month, min(l["emi_day"], calendar.monthrange(target_year, target_month)[1]))
                details = "Monthly EMI"
            else:
                # Unknown day: shown at the start of its month, saying so, rather than on a
                # made-up date. This month's stays listed (it may not have been debited yet).
                due_date = date(target_year, target_month, 1)
                details = "EMI day not set: add the EMI debit day on the loan"
                if m_offset == 0:
                    due_date = today
            if due_date >= today:
                events.append({
                    "title": f"{l['name']} EMI",
                    "amount": l["emi"],
                    "type": "EMI",
                    "date": due_date.isoformat(),
                    "details": details,
                })
                
        # B. Goals Savings contributions (1st of month)
        for g in goals:
            due_date = date(target_year, target_month, 1)
            if due_date >= today:
                events.append({
                    "title": f"{g['name']} Allocation",
                    "amount": g["monthly_contribution"],
                    "type": "SIP",
                    "date": due_date.isoformat(),
                    "details": "Monthly Goal Contribution"
                })
                
        # C. Insurance Premiums
        for ins in insurance:
            if not ins.get("renewal_date"):
                continue
            try:
                ren_d = datetime.strptime(ins["renewal_date"], "%Y-%m-%d").date()
                freq = ins["frequency"].lower()
                is_due = False
                
                months_diff = (target_year - ren_d.year) * 12 + (target_month - ren_d.month)
                
                if freq == "monthly":
                    is_due = True
                elif freq == "quarterly" and months_diff >= 0 and months_diff % 3 == 0:
                    is_due = True
                elif freq == "semi-annually" and months_diff >= 0 and months_diff % 6 == 0:
                    is_due = True
                elif freq in ("annually", "yearly", "annual") and months_diff >= 0 and months_diff % 12 == 0:
                    is_due = True
                    
                if is_due:
                    due_day = min(ren_d.day, calendar.monthrange(target_year, target_month)[1])
                    due_date = date(target_year, target_month, due_day)
                    if due_date >= today:
                        events.append({
                            "title": f"{ins['company']} Premium",
                            "amount": ins["premium"],
                            "type": "Insurance",
                            "date": due_date.isoformat(),
                            "details": f"Policy: {ins['policy_number']}"
                        })
            except Exception as e:
                print("Error computing insurance calendar date:", e)

    # D. Credit cards: once each, not every month (the balance was being repeated for all 12
    # months). What the last statement billed is due on the next due date; what's been spent
    # since is billed on the next statement date and due on the due date after that.
    def on_day_from(day, start):
        for add in range(3):
            y = start.year + (start.month - 1 + add) // 12
            m = (start.month - 1 + add) % 12 + 1
            d = date(y, m, min(day, calendar.monthrange(y, m)[1]))
            if d >= start:
                return d
        return start

    for cc in credit_cards:
        billed = min(cc["statement_balance"], cc["outstanding"])
        unbilled = max(cc["outstanding"] - billed, 0.0)
        if billed > 0:
            events.append({
                "title": f"{cc['name']}: bill due",
                "amount": billed,
                "type": "Credit Card",
                "date": on_day_from(cc["due_date"], today).isoformat(),
                "details": f"{cc['issuer']} · last statement",
            })
        if unbilled > 0:
            statement = on_day_from(cc["statement_date"], today)
            events.append({
                "title": f"{cc['name']}: next bill (estimate)",
                "amount": unbilled,
                "type": "Credit Card",
                "date": on_day_from(cc["due_date"], date.fromordinal(statement.toordinal() + 1)).isoformat(),
                "details": f"{cc['issuer']} · spent since the last statement, billed on {statement.strftime('%d %b')}",
            })

    try:
        horizon = date(today.year + 1, today.month, 1)
        events.extend(_tax_calendar_events(today, horizon))
    except Exception as e:
        print("[Calendar] Tax events unavailable:", e)
    events.sort(key=lambda x: x["date"])
    return events

# --- Multi-Loan Payoff Optimizer Endpoint ---
class CompareCardsRequest(BaseModel):
    owned_liability_ids: List[int] = []
    library_variant_ids: List[int] = []
    weights: Optional[dict] = None
    assumed_annual_spend: float = 600000.0

def _parse_card_json_fields(row: dict) -> dict:
    d = dict(row)
    for key in ("category_rules_json", "milestone_rules_json", "redemption_partners_json"):
        target = key.replace("_json", "")
        try:
            d[target] = json.loads(d.get(key) or "[]")
        except Exception:
            d[target] = []
    return d

@app.post("/api/credit-cards/compare")
def compare_credit_cards(req: CompareCardsRequest):
    total_selected = len(req.owned_liability_ids) + len(req.library_variant_ids)
    if total_selected < 2 or total_selected > 5:
        raise HTTPException(status_code=400, detail="Select between 2 and 5 cards to compare.")

    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        normalized = []
        if req.owned_liability_ids:
            placeholders = ", ".join("?" for _ in req.owned_liability_ids)
            cursor.execute(f"""
                SELECT l.id as liability_id, l.name, inst.name as lender_name, ccd.*
                FROM liabilities l
                JOIN credit_card_details ccd ON ccd.liability_id = l.id
                LEFT JOIN institutions inst ON l.lender_institution_id = inst.id
                WHERE l.id IN ({placeholders})
            """, req.owned_liability_ids)
            for row in cursor.fetchall():
                parsed = _parse_card_json_fields(dict(row))
                normalized.append(calc.normalize_card_for_comparison(parsed, "owned"))

        if req.library_variant_ids:
            placeholders = ", ".join("?" for _ in req.library_variant_ids)
            cursor.execute(f"SELECT * FROM card_variants WHERE id IN ({placeholders})", req.library_variant_ids)
            for row in cursor.fetchall():
                parsed = _parse_card_json_fields(dict(row))
                normalized.append(calc.normalize_card_for_comparison(parsed, "library"))

        if len(normalized) < 2:
            raise HTTPException(status_code=400, detail="Could not resolve at least 2 valid cards to compare.")

        return calc.compare_cards(normalized, req.weights, req.assumed_annual_spend)
    except HTTPException:
        raise
    finally:
        conn.close()

@app.get("/api/credit-cards/annual-review")
def get_annual_review(liability_id: Optional[int] = None):
    conn = get_db_connection()
    try:
        return generate_credit_card_annual_reviews(conn, liability_id)
    finally:
        conn.close()

class LoanOptimizerRequest(BaseModel):
    loans: list
    extra_monthly_budget: float
    prepayment_frequency: Optional[str] = "monthly"
    prepayment_amount: Optional[float] = 0.0
    prepayment_month: Optional[int] = 3
    prepayment_day: Optional[int] = 31

@app.post("/api/calculations/loan-optimizer")
def optimize_loans(req: LoanOptimizerRequest):
    try:
        baseline = calc.simulate_multi_loan_payoff(
            req.loans, 
            req.extra_monthly_budget, 
            "baseline",
            prepayment_frequency=req.prepayment_frequency,
            prepayment_amount=req.prepayment_amount,
            prepayment_month=req.prepayment_month,
            prepayment_day=req.prepayment_day
        )
        avalanche = calc.simulate_multi_loan_payoff(
            req.loans, 
            req.extra_monthly_budget, 
            "avalanche",
            prepayment_frequency=req.prepayment_frequency,
            prepayment_amount=req.prepayment_amount,
            prepayment_month=req.prepayment_month,
            prepayment_day=req.prepayment_day
        )
        snowball = calc.simulate_multi_loan_payoff(
            req.loans, 
            req.extra_monthly_budget, 
            "snowball",
            prepayment_frequency=req.prepayment_frequency,
            prepayment_amount=req.prepayment_amount,
            prepayment_month=req.prepayment_month,
            prepayment_day=req.prepayment_day
        )
        
        return {
            "baseline": baseline,
            "avalanche": avalanche,
            "snowball": snowball
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# --- Monthly Cash Flows Router ---
@app.get("/api/cash-flows")
def get_cash_flows():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM monthly_cash_flows ORDER BY month DESC")
        return [dict(row) for row in cursor.fetchall()]
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@app.post("/api/cash-flows")
def save_cash_flow(cf: CashFlowSchema):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT id FROM monthly_cash_flows WHERE month = ?", (cf.month,))
        existing = cursor.fetchone()
        
        if existing:
            cursor.execute("""
                UPDATE monthly_cash_flows
                SET net_salary = ?, other_income = ?, household_expenses = ?, emi = ?, 
                    insurance_premium = ?, investments = ?, one_time_expenses = ?, notes = ?
                WHERE id = ?
            """, (cf.net_salary, cf.other_income, cf.household_expenses, cf.emi, 
                  cf.insurance_premium, cf.investments, cf.one_time_expenses, cf.notes, existing["id"]))
        else:
            cursor.execute("""
                INSERT INTO monthly_cash_flows (month, net_salary, other_income, household_expenses, emi, 
                                                insurance_premium, investments, one_time_expenses, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (cf.month, cf.net_salary, cf.other_income, cf.household_expenses, cf.emi, 
                  cf.insurance_premium, cf.investments, cf.one_time_expenses, cf.notes))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@app.delete("/api/cash-flows/{cf_id}")
def delete_cash_flow(cf_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM monthly_cash_flows WHERE id = ?", (cf_id,))
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@app.post("/api/calculations/cash-flow/simulate")
def simulate_cash_flow(req: CashFlowSimulationRequest):
    try:
        # 1. Current State Calculations
        curr_income = req.net_salary + req.other_income
        curr_expenses = req.household_expenses + req.one_time_expenses
        curr_surplus = curr_income - curr_expenses - req.emi - req.insurance_premium - req.investments
        # Savings rate: the share of income not spent. Money invested is saved, so it
        # counts (as the target_savings_rate rule's Knowledge article defines it);
        # EMIs and premiums are commitments, so they don't.
        curr_savings_rate = ((curr_surplus + req.investments) / curr_income * 100.0) if curr_income > 0 else 0.0
        curr_debt_ratio = (req.emi / curr_income * 100.0) if curr_income > 0 else 0.0
        curr_investment_rate = (req.investments / curr_income * 100.0) if curr_income > 0 else 0.0
        curr_expense_ratio = (req.household_expenses / curr_income * 100.0) if curr_income > 0 else 0.0
        
        # 2. Simulated State Calculations
        sim_salary = req.net_salary * (1 + req.salary_change_pct / 100.0) + req.salary_change_abs
        sim_income = sim_salary + req.other_income
        
        sim_expenses = (req.household_expenses * (1 - req.expense_reduction_pct / 100.0) - req.expense_reduction_abs) + req.one_time_expenses
        sim_emi = max(req.emi + req.new_emi - req.emi_reduction, 0.0)
        sim_investments = req.investments + req.additional_investment
        
        sim_surplus = sim_income - sim_expenses - sim_emi - req.insurance_premium - sim_investments
        
        sim_savings_rate = ((sim_surplus + sim_investments) / sim_income * 100.0) if sim_income > 0 else 0.0
        sim_debt_ratio = (sim_emi / sim_income * 100.0) if sim_income > 0 else 0.0
        sim_investment_rate = (sim_investments / sim_income * 100.0) if sim_income > 0 else 0.0
        sim_expense_ratio = (sim_expenses / sim_income * 100.0) if sim_income > 0 else 0.0
        
        # 3. Financial Impact Statement
        impact_parts = []
        surplus_diff = sim_surplus - curr_surplus
        if surplus_diff != 0:
            direction = "increases" if surplus_diff > 0 else "decreases"
            impact_parts.append(f"Money left over each month {direction} by {inr(abs(surplus_diff))} (from {inr(curr_surplus)} to {inr(sim_surplus)}).")
            
        sr_diff = sim_savings_rate - curr_savings_rate
        if sr_diff != 0:
            direction = "grows" if sr_diff > 0 else "drops"
            impact_parts.append(f"Savings rate {direction} by {abs(sr_diff):.1f}% (from {curr_savings_rate:.1f}% to {sim_savings_rate:.1f}%).")
            
        dr_diff = sim_debt_ratio - curr_debt_ratio
        if dr_diff != 0:
            direction = "increases" if dr_diff > 0 else "decreases"
            impact_parts.append(f"Debt burden ratio {direction} by {abs(dr_diff):.1f}% (from {curr_debt_ratio:.1f}% to {sim_debt_ratio:.1f}%).")
            
        if req.one_time_bonus > 0:
            impact_parts.append(f"A one-time bonus of {inr(req.one_time_bonus)} adds to what you can invest.")
            
        impact_summary = " ".join(impact_parts) if impact_parts else "No cash flow impact observed."
        
        # 4. Recommended Actions
        recommendations = []
        if sim_surplus > curr_surplus:
            extra = sim_surplus - curr_surplus
            recommendations.append(f"Put the extra {inr(extra)} a month towards your SIPs or prepaying a high-interest loan.")
        if sim_debt_ratio > 40:
            recommendations.append("Warning: Debt burden ratio exceeds 40%. Consider deferring new loans and focusing on debt prepayment.")
        elif sim_savings_rate < 20:
            recommendations.append("Notice: Savings rate is below 20%. Look for areas to reduce discretionary household spending.")
        else:
            recommendations.append("Your cash flow parameters are healthy. Consider mapping extra savings to long-term goals.")
            
        if req.one_time_bonus > 0:
            recommendations.append(f"A common split for the {inr(req.one_time_bonus)} bonus: about 70% to loan prepayment or long-term investments, 30% kept as cash or for spending.")
            
        return {
            "current": {
                "income": curr_income,
                "expenses": curr_expenses,
                "emi": req.emi,
                "insurance": req.insurance_premium,
                "investments": req.investments,
                "surplus": curr_surplus,
                "savings_rate": round(curr_savings_rate, 2),
                "debt_ratio": round(curr_debt_ratio, 2),
                "investment_rate": round(curr_investment_rate, 2),
                "expense_ratio": round(curr_expense_ratio, 2)
            },
            "simulated": {
                "income": sim_income,
                "expenses": sim_expenses,
                "emi": sim_emi,
                "insurance": req.insurance_premium,
                "investments": sim_investments,
                "surplus": sim_surplus,
                "savings_rate": round(sim_savings_rate, 2),
                "debt_ratio": round(sim_debt_ratio, 2),
                "investment_rate": round(sim_investment_rate, 2),
                "expense_ratio": round(sim_expense_ratio, 2)
            },
            "impact_summary": impact_summary,
            "recommendations": recommendations
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    import uvicorn
    # init_db() runs in the lifespan, after the startup backup.
    uvicorn.run(app, host="127.0.0.1", port=8000)

